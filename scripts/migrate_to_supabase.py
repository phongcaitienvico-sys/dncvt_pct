#!/usr/bin/env python3
"""
Công cụ di trú & đối soát dữ liệu từ Google Sheets / JSON sang Supabase (PostgreSQL).

Hỗ trợ 2 định dạng nguồn ngày Cutover:
  1. JSON định dạng dev/data/db.json (tabs.<name>.headers & rows)
  2. Excel .xlsx xuất từ nút "Xuất toàn bộ ra Excel" (admin.html) hoặc out/db_khoi_tao.xlsx

Ba chế độ vận hành:
  python scripts/migrate_to_supabase.py dev/data/db.json --dry-run
  python scripts/migrate_to_supabase.py dev/data/db.json --sql-out out/cutover.sql
  python scripts/migrate_to_supabase.py dev/data/db.json --execute
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import secrets
import sys
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from collections import OrderedDict, defaultdict
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl

INITIAL_SEQ: dict[int, int] = {2026: 294}
MS_CODE_RE = re.compile(r"^[Mm][Ss]\d+$")
CTRL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
UUID_RE = re.compile(r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$")
TZ_SUFFIX_RE = re.compile(r"(?:[Zz]|[+-]\d{2}(?::?\d{2})?)$")

TIMESTAMPTZ_FIELDS: dict[str, tuple[str, ...]] = {
    "lines": ("created_at", "updated_at"),
    "requests": ("issued_at",),
    "receipts": ("created_at",),
    "audit_log": ("at",),
    "users": ("created_at",),
}

ALL_TABLES: tuple[str, ...] = (
    "lines",
    "requests",
    "receipts",
    "suppliers",
    "project_items",
    "material_groups",
    "units",
    "uoms",
    "signers",
    "users",
    "audit_log",
    "sequences",
)

# Thứ tự nạp vào PostgreSQL đảm bảo toàn vẹn khóa ngoại (requests trước lines, lines trước receipts)
SQL_ID_TABLES: tuple[str, ...] = (
    "suppliers",
    "project_items",
    "material_groups",
    "units",
    "uoms",
    "signers",
    "requests",
    "lines",
    "receipts",
    "audit_log",
    "sequences",
)

TABLE_COLUMNS: dict[str, list[str]] = {
    "suppliers": ["id", "code", "name", "address", "phone", "email", "tax_code", "field"],
    "project_items": ["id", "project_code", "project_name", "item_code", "item_name", "owner", "status"],
    "material_groups": ["id", "name", "description"],
    "units": ["id", "name"],
    "uoms": ["id", "name"],
    "signers": ["id", "unit_contains", "position", "title", "name"],
    "requests": ["id", "code", "year", "seq", "issued_at", "issued_by", "header_json", "line_count", "legacy_flag"],
    "lines": [
        "id", "request_code", "unit", "material_group", "name", "uom",
        "qty_requested", "qty_on_hand", "need_date", "project_code",
        "project_name", "purpose", "equipment_name", "delivery_location",
        "requester", "request_date", "supplier", "unit_price", "note",
        "legacy_flag", "created_by", "created_at", "updated_at",
        "qty_received", "last_received_date",
    ],
    "receipts": ["id", "line_id", "received_date", "qty", "received_by", "note", "created_at", "created_by"],
    "audit_log": ["id", "at", "user", "table", "row_id", "field", "old", "new"],
    "sequences": ["id", "key", "last_value"],
}

# Đồng bộ chuẩn từ apps-script/Schema.js
DEFAULT_SIGNERS: list[dict[str, str]] = [
    {"unit_contains": "", "position": "1", "title": "Người lập phiếu", "name": ""},
    {"unit_contains": "", "position": "2", "title": "Người xem xét", "name": "Vũ Văn Trường"},
    {"unit_contains": "", "position": "3", "title": "T.P Thiết bị", "name": "Vũ Văn Trường"},
    {"unit_contains": "", "position": "4", "title": "Giám đốc duyệt", "name": "Phạm Danh Mạnh"},
    {"unit_contains": "Cải tiến", "position": "1", "title": "Người lập phiếu", "name": ""},
    {"unit_contains": "Cải tiến", "position": "2", "title": "Người xem xét", "name": "Trần Việt Tiệp"},
    {"unit_contains": "Cải tiến", "position": "3", "title": "T.P Cải tiến", "name": "Trần Việt Tiệp"},
    {"unit_contains": "Cải tiến", "position": "4", "title": "Giám đốc duyệt", "name": "Phạm Danh Mạnh"},
]


@dataclass
class MigrationBundle:
    tables: dict[str, list[dict[str, Any]]]
    merged_project_items: list[dict[str, Any]] = field(default_factory=list)
    stats: dict[str, Any] = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Hàm tiện ích chuẩn hóa kiểu dữ liệu
# ---------------------------------------------------------------------------


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def num_out(x: float | int | None) -> int | float | None:
    if x is None:
        return None
    fx = float(x)
    return int(fx) if fx.is_integer() else fx


def parse_int(v: Any, field_name: str = "int") -> int | None:
    if _is_blank(v):
        return None
    if isinstance(v, bool):
        raise ValueError(f"Trường {field_name} không hợp lệ (bool): {v!r}")
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        if v.is_integer():
            return int(v)
        raise ValueError(f"Trường {field_name} phải là số nguyên: {v!r}")
    s = str(v).strip()
    try:
        fv = float(s)
    except ValueError as exc:
        raise ValueError(f"Trường {field_name} không phải số nguyên: {v!r}") from exc
    if not fv.is_integer():
        raise ValueError(f"Trường {field_name} phải là số nguyên: {v!r}")
    return int(fv)


def parse_num(v: Any, field_name: str = "num") -> int | float | None:
    if _is_blank(v):
        return None
    if isinstance(v, bool):
        raise ValueError(f"Trường {field_name} không hợp lệ (bool): {v!r}")
    if isinstance(v, (int, float)):
        return num_out(v)
    s = str(v).strip().replace(",", ".")
    try:
        return num_out(float(s))
    except ValueError as exc:
        raise ValueError(f"Trường {field_name} không phải số: {v!r}") from exc


def parse_bool(v: Any) -> bool:
    if isinstance(v, bool):
        return v
    if v is None:
        return False
    return str(v).strip().upper() in ("TRUE", "1")


def parse_date_str(v: Any, field_name: str = "date") -> str | None:
    if _is_blank(v):
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return dt.date(int(m[1]), int(m[2]), int(m[3])).isoformat()
    m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})$", s)
    if m:
        return dt.date(int(m[3]), int(m[2]), int(m[1])).isoformat()
    raise ValueError(f"Trường {field_name} không phải ngày hợp lệ: {v!r}")


def parse_datetime_str(v: Any) -> str | None:
    if _is_blank(v):
        return None
    if isinstance(v, dt.datetime):
        if v.tzinfo is not None:
            return v.isoformat(timespec="seconds")
        return v.strftime("%Y-%m-%dT%H:%M:%S")
    if isinstance(v, dt.date):
        return f"{v.isoformat()}T00:00:00"
    s = str(v).strip()
    return s or None


def ensure_vn_tz(ts: str | None) -> str | None:
    """Gắn hậu tố múi giờ Việt Nam (+07:00) cho chuỗi ISO naive để chống lệch +7h khi phiên DB ở UTC."""
    if _is_blank(ts):
        return None
    s = str(ts).strip()
    if TZ_SUFFIX_RE.search(s):
        return s
    return f"{s}+07:00"


def parse_user_id(v: Any) -> int | str:
    """Đọc users.id hỗ trợ cả số nguyên lịch sử (db.json) và UUID chuỗi (xuất từ Supabase rpc_export_all)."""
    if _is_blank(v):
        raise ValueError("Trường users.id không được để trống")
    if isinstance(v, bool):
        raise ValueError(f"Trường users.id không hợp lệ (bool): {v!r}")
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    s = clean_nfc_str(v)
    if not s:
        raise ValueError("Trường users.id không được để trống")
    if UUID_RE.fullmatch(s):
        return s.lower()
    try:
        fv = float(s)
        if fv.is_integer():
            return int(fv)
    except ValueError:
        pass
    return s


def parse_optional_uuid(v: Any) -> str | None:
    """Chỉ giữ lại chuỗi UUID hợp lệ (cho cột khóa ngoại kiểu uuid như receipts.created_by)."""
    s = clean_nfc_str(v)
    if not s:
        return None
    return s.lower() if UUID_RE.fullmatch(s) else None


def clean_nfc_str(v: Any, *, empty_as_str: bool = False, collapse_spaces: bool = False) -> str | None:
    if v is None:
        return "" if empty_as_str else None
    if isinstance(v, bool):
        s = "TRUE" if v else ""
    elif isinstance(v, float) and v.is_integer():
        s = str(int(v))
    elif isinstance(v, (dt.datetime, dt.date)):
        s = v.isoformat()
    else:
        s = str(v)
    s = CTRL_CHARS_RE.sub("", unicodedata.normalize("NFC", s)).strip()
    if collapse_spaces:
        s = re.sub(r"\s+", " ", s)
    if not s:
        return "" if empty_as_str else None
    return s


def catalog_key(v: Any) -> str:
    s = clean_nfc_str(v, collapse_spaces=True)
    return s.lower() if s else ""


def normalize_ms_code(code: str | None) -> tuple[str | None, bool]:
    """Chuẩn hóa mã dạng ^[Mm][Ss]\\d+$ về chữ IN HOA. Trả (mã_mới, đã_đổi_hoa_thường)."""
    if not code:
        return None, False
    if MS_CODE_RE.fullmatch(code):
        upper = code.upper()
        return upper, upper != code
    return code, False


# ---------------------------------------------------------------------------
# 1. Đọc dữ liệu nguồn (.json hoặc .xlsx)
# ---------------------------------------------------------------------------


def load_source(path: Path | str) -> dict[str, list[dict[str, Any]]]:
    src = Path(path)
    if not src.exists():
        raise FileNotFoundError(f"Không tìm thấy file nguồn: {src}")

    suffix = src.suffix.lower()
    if suffix == ".json":
        return _load_json_source(src)
    if suffix in (".xlsx", ".xlsm"):
        return _load_xlsx_source(src)
    raise ValueError(f"Định dạng file không hỗ trợ ({suffix}): chỉ nhận .json hoặc .xlsx")


def _load_json_source(src: Path) -> dict[str, list[dict[str, Any]]]:
    payload = json.loads(src.read_text(encoding="utf-8"))
    tabs = payload.get("tabs") if isinstance(payload, dict) else None
    if isinstance(tabs, dict):
        out: dict[str, list[dict[str, Any]]] = {}
        for name in ALL_TABLES:
            tab = tabs.get(name)
            if not tab:
                out[name] = []
                continue
            headers = [str(h).strip() for h in tab.get("headers", [])]
            id_idx = headers.index("id") if "id" in headers else -1
            rows_out: list[dict[str, Any]] = []
            for raw_row in tab.get("rows", []):
                if id_idx >= 0:
                    id_val = raw_row[id_idx] if id_idx < len(raw_row) else None
                    if _is_blank(id_val):
                        continue
                elif all(_is_blank(c) for c in raw_row):
                    continue
                padded = list(raw_row) + [None] * max(0, len(headers) - len(raw_row))
                rows_out.append(dict(zip(headers, padded)))
            out[name] = rows_out
        return out

    # Hỗ trợ JSON phẳng từ rpc_export_all hoặc bản ghi private.backup_snapshots (bọc trong 'payload' / 'data')
    flat: Any = payload
    if isinstance(payload, dict) and not any(isinstance(payload.get(t), list) for t in ALL_TABLES):
        if isinstance(payload.get("payload"), dict):
            flat = payload["payload"]
        elif isinstance(payload.get("data"), dict):
            flat = payload["data"]

    if not isinstance(flat, dict) or not any(isinstance(flat.get(t), list) for t in ALL_TABLES):
        raise ValueError(
            f"File JSON không đúng cấu trúc db.json (thiếu 'tabs') hoặc JSON phẳng rpc_export_all: {src}"
        )

    out = {}
    for name in ALL_TABLES:
        raw_list = flat.get(name)
        if not raw_list:
            out[name] = []
            continue
        if not isinstance(raw_list, list):
            raise ValueError(f"Bảng '{name}' trong JSON phẳng không phải danh sách (list): {src}")
        rows_out = []
        for item in raw_list:
            if not isinstance(item, dict):
                continue
            if "id" in item and _is_blank(item.get("id")):
                continue
            if all(_is_blank(v) for v in item.values()):
                continue
            rows_out.append(dict(item))
        out[name] = rows_out
    return out


def _load_xlsx_source(src: Path) -> dict[str, list[dict[str, Any]]]:
    wb_val = openpyxl.load_workbook(src, read_only=True, data_only=True)
    wb_raw = openpyxl.load_workbook(src, read_only=True, data_only=False)
    try:
        sheet_val = {ws.title: ws for ws in wb_val.worksheets}
        sheet_raw = {ws.title: ws for ws in wb_raw.worksheets}
        out: dict[str, list[dict[str, Any]]] = {}
        for name in ALL_TABLES:
            ws_v = sheet_val.get(name)
            ws_r = sheet_raw.get(name)
            if ws_v is None or ws_r is None:
                out[name] = []
                continue
            it_v = ws_v.iter_rows(values_only=True)
            it_r = ws_r.iter_rows(values_only=True)
            first_row = next(it_v, None)
            next(it_r, None)
            if not first_row:
                out[name] = []
                continue
            headers = [str(h).strip() if h is not None else "" for h in first_row]
            id_idx = headers.index("id") if "id" in headers else -1
            rows_out: list[dict[str, Any]] = []
            for row_v, row_r in zip(it_v, it_r):
                if row_v is None:
                    continue
                raw_row = [
                    (cv if cv is not None else cr)
                    for cv, cr in zip(row_v, row_r)
                ]
                if id_idx >= 0:
                    id_val = raw_row[id_idx] if id_idx < len(raw_row) else None
                    if _is_blank(id_val):
                        continue
                elif all(_is_blank(c) for c in raw_row):
                    continue
                padded = list(raw_row) + [None] * max(0, len(headers) - len(raw_row))
                rows_out.append({h: v for h, v in zip(headers, padded) if h})
            out[name] = rows_out
        return out
    finally:
        wb_val.close()
        wb_raw.close()


# ---------------------------------------------------------------------------
# 2. Làm sạch & chuẩn hóa 8 dị thường dữ liệu lịch sử
# ---------------------------------------------------------------------------


def transform_data(raw: dict[str, list[dict[str, Any]]]) -> MigrationBundle:
    tables: dict[str, list[dict[str, Any]]] = {}
    stats: dict[str, Any] = {
        "counts_before": {t: len(raw.get(t, [])) for t in ALL_TABLES},
        "project_items_ms_normalized": 0,
        "lines_ms_normalized": 0,
        "lines_nfc_fixed": 0,
        "seeded_signers": False,
    }

    # --- 2.1 project_items: khử trùng khóa project_code, giữ id nhỏ hơn & gộp cột đầy đủ nhất, Ms->MS ---
    pi_groups: OrderedDict[tuple[str, Any], list[dict[str, Any]]] = OrderedDict()
    for r in raw.get("project_items", []):
        rid = parse_int(r.get("id"), "project_items.id")
        if rid is None:
            raise ValueError("project_items có dòng thiếu id")
        cleaned = {
            "id": rid,
            "project_code": clean_nfc_str(r.get("project_code"), collapse_spaces=True),
            "project_name": clean_nfc_str(r.get("project_name")),
            "item_code": clean_nfc_str(r.get("item_code"), collapse_spaces=True),
            "item_name": clean_nfc_str(r.get("item_name")),
            "owner": clean_nfc_str(r.get("owner")),
            "status": clean_nfc_str(r.get("status")),
        }
        key_str = catalog_key(cleaned["project_code"])
        grp_key = ("pc", key_str) if key_str else ("id", rid)
        pi_groups.setdefault(grp_key, []).append(cleaned)

    merged_info: list[dict[str, Any]] = []
    pi_out: list[dict[str, Any]] = []
    for _, grp in pi_groups.items():
        grp_sorted = sorted(grp, key=lambda x: x["id"])
        kept_id = grp_sorted[0]["id"]
        merged_row: dict[str, Any] = {"id": kept_id}
        for col in ("project_code", "project_name", "item_code", "item_name", "owner", "status"):
            non_empty = [x[col] for x in grp_sorted if x[col]]
            # max(..., key=len) giữ phần tử đầu tiên (id nhỏ nhất) nếu độ dài bằng nhau
            merged_row[col] = max(non_empty, key=len) if non_empty else None

        norm_pc, changed_pc = normalize_ms_code(merged_row["project_code"])
        if changed_pc:
            stats["project_items_ms_normalized"] += 1
        merged_row["project_code"] = norm_pc

        norm_ic, _ = normalize_ms_code(merged_row["item_code"])
        merged_row["item_code"] = norm_ic

        if len(grp_sorted) > 1:
            merged_info.append(
                {
                    "project_code": merged_row["project_code"],
                    "kept_id": kept_id,
                    "dropped_ids": [x["id"] for x in grp_sorted[1:]],
                    "raw_codes": [x["project_code"] for x in grp_sorted],
                }
            )
        pi_out.append(merged_row)

    pi_out.sort(key=lambda x: x["id"])
    tables["project_items"] = pi_out

    # --- 2.2 receipts: giữ nguyên lịch sử, cho phép received_date NULL, created_by NULL ---
    receipts_out: list[dict[str, Any]] = []
    receipt_agg: dict[int, list[Any]] = {}  # line_id -> [sum_qty, max_received_date]
    for r in raw.get("receipts", []):
        rid = parse_int(r.get("id"), "receipts.id")
        line_id = parse_int(r.get("line_id"), "receipts.line_id")
        qty = parse_num(r.get("qty"), "receipts.qty")
        if rid is None or line_id is None or qty is None:
            raise ValueError(f"receipts dòng không hợp lệ: {r!r}")
        rdate = parse_date_str(r.get("received_date"), "receipts.received_date")
        rec_row = {
            "id": rid,
            "line_id": line_id,
            "received_date": rdate,
            "qty": qty,
            "received_by": clean_nfc_str(r.get("received_by")),
            "note": clean_nfc_str(r.get("note")),
            "created_at": parse_datetime_str(r.get("created_at")),
            "created_by": parse_optional_uuid(r.get("created_by")),
        }
        receipts_out.append(rec_row)

        agg = receipt_agg.setdefault(line_id, [0.0, None])
        agg[0] += float(qty)
        if rdate and (agg[1] is None or rdate > agg[1]):
            agg[1] = rdate

    tables["receipts"] = receipts_out

    # --- 2.3 lines: giữ nguyên id & created_at/updated_at, Ms->MS, NFC, tính qty_received & last_received_date ---
    lines_out: list[dict[str, Any]] = []
    line_str_fields = (
        "request_code", "unit", "material_group", "name", "uom",
        "project_name", "purpose", "equipment_name", "delivery_location",
        "requester", "supplier", "note", "created_by",
    )
    for r in raw.get("lines", []):
        lid = parse_int(r.get("id"), "lines.id")
        if lid is None:
            raise ValueError("lines có dòng thiếu id")

        for sf in (*line_str_fields, "project_code"):
            rv = r.get(sf)
            if isinstance(rv, str) and not unicodedata.is_normalized("NFC", rv):
                stats["lines_nfc_fixed"] += 1
                break

        raw_pc = clean_nfc_str(r.get("project_code"), collapse_spaces=True)
        norm_pc, changed_pc = normalize_ms_code(raw_pc)
        if changed_pc:
            stats["lines_ms_normalized"] += 1

        agg = receipt_agg.get(lid)
        qty_recv = num_out(round(agg[0], 3)) if agg else 0
        last_rdate = agg[1] if agg else None

        line_row: dict[str, Any] = {
            "id": lid,
            "request_code": clean_nfc_str(r.get("request_code")),
            "unit": clean_nfc_str(r.get("unit")),
            "material_group": clean_nfc_str(r.get("material_group")),
            "name": clean_nfc_str(r.get("name")),
            "uom": clean_nfc_str(r.get("uom")),
            "qty_requested": parse_num(r.get("qty_requested"), "lines.qty_requested"),
            "qty_on_hand": parse_num(r.get("qty_on_hand"), "lines.qty_on_hand"),
            "need_date": parse_date_str(r.get("need_date"), "lines.need_date"),
            "project_code": norm_pc,
            "project_name": clean_nfc_str(r.get("project_name")),
            "purpose": clean_nfc_str(r.get("purpose")),
            "equipment_name": clean_nfc_str(r.get("equipment_name")),
            "delivery_location": clean_nfc_str(r.get("delivery_location")),
            "requester": clean_nfc_str(r.get("requester")),
            "request_date": parse_date_str(r.get("request_date"), "lines.request_date"),
            "supplier": clean_nfc_str(r.get("supplier")),
            "unit_price": parse_num(r.get("unit_price"), "lines.unit_price"),
            "note": clean_nfc_str(r.get("note")),
            "legacy_flag": parse_bool(r.get("legacy_flag")),
            "created_by": clean_nfc_str(r.get("created_by")),
            "created_at": parse_datetime_str(r.get("created_at")),
            "updated_at": parse_datetime_str(r.get("updated_at")),
            "qty_received": qty_recv,
            "last_received_date": last_rdate,
        }
        lines_out.append(line_row)

    tables["lines"] = lines_out

    # --- 2.4 requests: BẤT BIẾN TUYỆT ĐỐI header_json, cho phép issued_at NULL ---
    requests_out: list[dict[str, Any]] = []
    for r in raw.get("requests", []):
        rid = parse_int(r.get("id"), "requests.id")
        code = clean_nfc_str(r.get("code"))
        year = parse_int(r.get("year"), "requests.year")
        seq = parse_int(r.get("seq"), "requests.seq")
        if rid is None or not code or year is None or seq is None:
            raise ValueError(f"requests dòng thiếu thông tin bắt buộc: {r!r}")

        raw_hdr = r.get("header_json")
        if isinstance(raw_hdr, str):
            try:
                parsed_hdr = json.loads(raw_hdr)
            except Exception as exc:
                raise ValueError(f"requests id={rid} ({code}) có header_json không phải JSON hợp lệ: {exc}") from exc
            if not isinstance(parsed_hdr, dict):
                raise ValueError(f"requests id={rid} ({code}) có header_json không phải JSON object")
            immutable_hdr = raw_hdr  # Giữ nguyên từng ký tự, không sửa đổi
        elif isinstance(raw_hdr, dict):
            immutable_hdr = json.dumps(raw_hdr, ensure_ascii=False)
        else:
            raise ValueError(f"requests id={rid} ({code}) thiếu header_json")

        requests_out.append(
            {
                "id": rid,
                "code": code,
                "year": year,
                "seq": seq,
                "issued_at": parse_datetime_str(r.get("issued_at")),
                "issued_by": clean_nfc_str(r.get("issued_by")),
                "header_json": immutable_hdr,
                "line_count": parse_int(r.get("line_count"), "requests.line_count") or 0,
                "legacy_flag": parse_bool(r.get("legacy_flag")),
            }
        )
    tables["requests"] = requests_out

    # --- 2.5 Các bảng danh mục còn lại ---
    tables["suppliers"] = [
        {
            "id": parse_int(r.get("id"), "suppliers.id"),
            "code": clean_nfc_str(r.get("code")),
            "name": clean_nfc_str(r.get("name")),
            "address": clean_nfc_str(r.get("address")),
            "phone": clean_nfc_str(r.get("phone")),
            "email": clean_nfc_str(r.get("email")),
            "tax_code": clean_nfc_str(r.get("tax_code")),
            "field": clean_nfc_str(r.get("field")),
        }
        for r in raw.get("suppliers", [])
    ]
    tables["material_groups"] = [
        {
            "id": parse_int(r.get("id"), "material_groups.id"),
            "name": clean_nfc_str(r.get("name")),
            "description": clean_nfc_str(r.get("description")),
        }
        for r in raw.get("material_groups", [])
    ]
    tables["units"] = [
        {"id": parse_int(r.get("id"), "units.id"), "name": clean_nfc_str(r.get("name"))}
        for r in raw.get("units", [])
    ]
    tables["uoms"] = [
        {"id": parse_int(r.get("id"), "uoms.id"), "name": clean_nfc_str(r.get("name"))}
        for r in raw.get("uoms", [])
    ]

    # --- 2.6 signers: nếu trống hoặc chưa có thì seed 8 dòng DEFAULT_SIGNERS ---
    raw_signers = raw.get("signers", [])
    if not raw_signers:
        stats["seeded_signers"] = True
        tables["signers"] = [
            {
                "id": idx,
                "unit_contains": s["unit_contains"],
                "position": s["position"],
                "title": s["title"],
                "name": s["name"],
            }
            for idx, s in enumerate(DEFAULT_SIGNERS, 1)
        ]
    else:
        tables["signers"] = [
            {
                "id": parse_int(r.get("id"), "signers.id"),
                "unit_contains": clean_nfc_str(r.get("unit_contains"), empty_as_str=True),
                "position": clean_nfc_str(r.get("position"), empty_as_str=True),
                "title": clean_nfc_str(r.get("title")),
                "name": clean_nfc_str(r.get("name"), empty_as_str=True),
            }
            for r in raw_signers
        ]

    # --- 2.7 users & audit_log ---
    tables["users"] = [
        {
            "id": parse_user_id(r.get("id")),
            "username": clean_nfc_str(r.get("username")),
            "full_name": clean_nfc_str(r.get("full_name")) or clean_nfc_str(r.get("username")),
            "role": "admin" if clean_nfc_str(r.get("role")) == "admin" else "user",
            "password_hash": clean_nfc_str(r.get("password_hash")),
            "salt": clean_nfc_str(r.get("salt")),
            "active": parse_bool(r.get("active")) if "active" in r else True,
            "created_at": parse_datetime_str(r.get("created_at")),
        }
        for r in raw.get("users", [])
    ]

    tables["audit_log"] = [
        {
            "id": parse_int(r.get("id"), "audit_log.id"),
            "at": parse_datetime_str(r.get("at")),
            "user": clean_nfc_str(r.get("user")),
            "table": clean_nfc_str(r.get("table")),
            "row_id": clean_nfc_str(
                r.get("row_id") if not _is_blank(r.get("row_id")) else r.get("record_id")
            ),
            "field": clean_nfc_str(r.get("field")),
            "old": clean_nfc_str(r.get("old")),
            "new": clean_nfc_str(r.get("new")),
        }
        for r in raw.get("audit_log", [])
    ]

    # --- 2.8 sequences: đảm bảo YC<year> >= MAX(seq) (legacy_flag=False) và YC2026 >= 294 ---
    required_yc: dict[int, int] = defaultdict(int)
    for req in requests_out:
        if not req["legacy_flag"]:
            required_yc[req["year"]] = max(required_yc[req["year"]], req["seq"])
    for yr, min_seq in INITIAL_SEQ.items():
        required_yc[yr] = max(required_yc[yr], min_seq)

    seq_out: list[dict[str, Any]] = []
    seen_keys: dict[str, dict[str, Any]] = {}
    max_seq_id = 0
    for r in raw.get("sequences", []):
        sid = parse_int(r.get("id"), "sequences.id")
        k = clean_nfc_str(r.get("key"))
        lv = parse_int(r.get("last_value"), "sequences.last_value") or 0
        if sid is None or not k:
            raise ValueError(f"sequences dòng không hợp lệ: {r!r}")
        max_seq_id = max(max_seq_id, sid)
        m = re.fullmatch(r"YC(\d{4})", k)
        if m:
            yr = int(m[1])
            lv = max(lv, required_yc.get(yr, 0))
        row_obj = {"id": sid, "key": k, "last_value": lv}
        seq_out.append(row_obj)
        seen_keys[k] = row_obj

    for yr in sorted(required_yc):
        k = f"YC{yr}"
        if k not in seen_keys:
            max_seq_id += 1
            seq_out.append({"id": max_seq_id, "key": k, "last_value": required_yc[yr]})

    tables["sequences"] = seq_out
    stats["counts_after"] = {t: len(tables[t]) for t in ALL_TABLES}
    return MigrationBundle(tables=tables, merged_project_items=merged_info, stats=stats)


# ---------------------------------------------------------------------------
# 3. Đối soát tự động cứng (sai 1 đồng hoặc thiếu 1 dòng là raise ValueError)
# ---------------------------------------------------------------------------


def reconcile(raw: dict[str, list[dict[str, Any]]], bundle: MigrationBundle) -> dict[str, Any]:
    raw_lines = raw.get("lines", [])
    after_lines = bundle.tables["lines"]
    if len(after_lines) != len(raw_lines):
        raise ValueError(f"Đối soát lines thất bại: số dòng sau ({len(after_lines)}) != trước ({len(raw_lines)})")

    before_line_ids = [parse_int(r.get("id"), "lines.id") for r in raw_lines]
    after_line_ids = [r["id"] for r in after_lines]
    if len(set(after_line_ids)) != len(after_lines) or set(after_line_ids) != set(before_line_ids):
        raise ValueError("Đối soát lines thất bại: tập id (STT) bị thay đổi hoặc trùng lặp")

    before_line_ts = {
        parse_int(r.get("id"), "lines.id"): (
            ensure_vn_tz(parse_datetime_str(r.get("created_at"))),
            ensure_vn_tz(parse_datetime_str(r.get("updated_at"))),
        )
        for r in raw_lines
    }
    for l_row in after_lines:
        after_ts = (
            ensure_vn_tz(l_row["created_at"]),
            ensure_vn_tz(l_row["updated_at"]),
        )
        if after_ts != before_line_ts[l_row["id"]]:
            raise ValueError(
                f"Đối soát lines thất bại: dòng id={l_row['id']} bị thay đổi created_at/updated_at lịch sử"
            )
        pc = l_row.get("project_code")
        if pc and MS_CODE_RE.fullmatch(pc) and pc != pc.upper():
            raise ValueError(f"Đối soát lines thất bại: còn mã project_code chưa IN HOA ({pc})")

    # Đối soát thành tiền SUM((qty_requested or 0) * (unit_price or 0))
    amount_before = sum(
        float(parse_num(r.get("qty_requested")) or 0) * float(parse_num(r.get("unit_price")) or 0)
        for r in raw_lines
    )
    amount_after = sum(
        float(r.get("qty_requested") or 0) * float(r.get("unit_price") or 0)
        for r in after_lines
    )
    amount_diff = round(amount_after - amount_before, 6)
    if amount_diff != 0:
        raise ValueError(
            f"Đối soát tổng thành tiền thất bại: trước={amount_before:,.6f}, sau={amount_after:,.6f}, lệch={amount_diff:,.6f}"
        )

    # Đối soát requests & tính bất biến của header_json
    raw_reqs = raw.get("requests", [])
    after_reqs = bundle.tables["requests"]
    if len(after_reqs) != len(raw_reqs):
        raise ValueError(f"Đối soát requests thất bại: số phiếu sau ({len(after_reqs)}) != trước ({len(raw_reqs)})")

    before_req_map = {parse_int(r.get("id"), "requests.id"): r for r in raw_reqs}
    if set(r["id"] for r in after_reqs) != set(before_req_map.keys()):
        raise ValueError("Đối soát requests thất bại: tập id phiếu không khớp")

    for req in after_reqs:
        orig = before_req_map[req["id"]]
        orig_hdr = orig.get("header_json")
        expected_hdr = orig_hdr if isinstance(orig_hdr, str) else json.dumps(orig_hdr, ensure_ascii=False)
        if req["header_json"] != expected_hdr:
            raise ValueError(f"Đối soát requests thất bại: header_json của phiếu {req['code']} bị biến đổi!")
        # Xác thực parse JSON lần nữa
        json.loads(req["header_json"])

    # Đối soát receipts
    raw_recs = raw.get("receipts", [])
    after_recs = bundle.tables["receipts"]
    if len(after_recs) != len(raw_recs):
        raise ValueError(f"Đối soát receipts thất bại: số dòng sau ({len(after_recs)}) != trước ({len(raw_recs)})")

    if {r["id"] for r in after_recs} != {parse_int(r.get("id"), "receipts.id") for r in raw_recs}:
        raise ValueError("Đối soát receipts thất bại: tập id nhận hàng không khớp")

    rec_sum_before = round(sum(float(parse_num(r.get("qty")) or 0) for r in raw_recs), 3)
    rec_sum_after = round(sum(float(r.get("qty") or 0) for r in after_recs), 3)
    rec_diff = round(rec_sum_after - rec_sum_before, 3)
    if rec_diff != 0:
        raise ValueError(
            f"Đối soát receipts thất bại: tổng qty trước={rec_sum_before}, sau={rec_sum_after}, lệch={rec_diff}"
        )

    # Đối soát project_items: sau + số dòng gộp == trước, 0 mã trùng còn sót
    raw_pi = raw.get("project_items", [])
    after_pi = bundle.tables["project_items"]
    extra_merged = sum(len(m["dropped_ids"]) for m in bundle.merged_project_items)
    if len(after_pi) + extra_merged != len(raw_pi):
        raise ValueError(
            f"Đối soát project_items thất bại: sau ({len(after_pi)}) + đã gộp ({extra_merged}) != trước ({len(raw_pi)})"
        )

    seen_pi_keys: set[str] = set()
    for p_row in after_pi:
        pc = p_row.get("project_code")
        if pc:
            if MS_CODE_RE.fullmatch(pc) and pc != pc.upper():
                raise ValueError(f"Đối soát project_items thất bại: còn mã Ms viết thường ({pc})")
            ck = catalog_key(pc)
            if ck in seen_pi_keys:
                raise ValueError(f"Đối soát project_items thất bại: còn sót mã trùng ({pc})")
            seen_pi_keys.add(ck)

    # Đối soát sequences
    seq_map = {s["key"]: s["last_value"] for s in bundle.tables["sequences"]}
    for req in after_reqs:
        if not req["legacy_flag"]:
            k = f"YC{req['year']}"
            if seq_map.get(k, -1) < req["seq"]:
                raise ValueError(f"Đối soát sequences thất bại: {k}={seq_map.get(k)} < seq {req['seq']}")
    for yr, min_seq in INITIAL_SEQ.items():
        if seq_map.get(f"YC{yr}", -1) < min_seq:
            raise ValueError(f"Đối soát sequences thất bại: YC{yr}={seq_map.get(f'YC{yr}')} < {min_seq}")

    summary = {
        "status": "PASS",
        "amount_before": amount_before,
        "amount_after": amount_after,
        "amount_diff": 0 if amount_diff == 0 else amount_diff,
        "receipt_qty_before": rec_sum_before,
        "receipt_qty_after": rec_sum_after,
        "receipt_qty_diff": 0 if rec_diff == 0 else rec_diff,
        "project_items_merged_pairs": len(bundle.merged_project_items),
        "project_items_dropped_rows": extra_merged,
    }
    bundle.stats["reconcile"] = summary
    return summary


# ---------------------------------------------------------------------------
# 4. Báo cáo đối soát Tiếng Việt
# ---------------------------------------------------------------------------


def format_report(bundle: MigrationBundle, source_name: str = "") -> str:
    rec = bundle.stats.get("reconcile", {})
    status = rec.get("status", "PASS")
    before = bundle.stats.get("counts_before", {})
    after = bundle.stats.get("counts_after", {})

    lines = [
        f"BÁO CÁO ĐỐI SOÁT DI TRÚ SUPABASE{f' — {source_name}' if source_name else ''}",
        "=" * 72,
        f"Trạng thái đối soát tổng thể     : {status}",
        "",
        "1. Số dòng từng bảng (Trước -> Sau chuẩn hóa):",
    ]
    for tbl in ALL_TABLES:
        note = ""
        if tbl == "project_items" and bundle.merged_project_items:
            dropped = sum(len(m["dropped_ids"]) for m in bundle.merged_project_items)
            note = f"  (đã gộp {dropped} dòng trùng từ {len(bundle.merged_project_items)} mã)"
        elif tbl == "signers" and bundle.stats.get("seeded_signers"):
            note = "  (tự động khởi tạo 8 dòng DEFAULT_SIGNERS)"
        lines.append(f"   - {tbl:16s}: {before.get(tbl, 0):6,d} -> {after.get(tbl, 0):6,d}{note}")

    lines += [
        "",
        "2. Kết quả làm sạch & chuẩn hóa dị thường lịch sử:",
        f"   - Số cặp project_items trùng đã gộp : {len(bundle.merged_project_items):,}",
        f"   - Mã Ms->MS chuẩn hóa (project_items): {bundle.stats.get('project_items_ms_normalized', 0):,}",
        f"   - Mã Ms->MS chuẩn hóa (lines)       : {bundle.stats.get('lines_ms_normalized', 0):,}",
        f"   - Dòng lines chuẩn hóa Unicode NFC  : {bundle.stats.get('lines_nfc_fixed', 0):,}",
        f"   - Phiếu lịch sử trống issued_at     : {sum(1 for r in bundle.tables['requests'] if r['issued_at'] is None):,}",
        f"   - Lần nhận hàng trống received_date : {sum(1 for r in bundle.tables['receipts'] if r['received_date'] is None):,}",
    ]

    if bundle.merged_project_items:
        lines.append("")
        lines.append("   Danh sách chi tiết các cặp project_items trùng đã gộp:")
        for m in bundle.merged_project_items:
            dropped_str = ", ".join(str(i) for i in m["dropped_ids"])
            raw_str = " / ".join(str(c) for c in m["raw_codes"])
            lines.append(f"     * {m['project_code']} (gốc: {raw_str}): giữ id={m['kept_id']}, gộp bỏ id=[{dropped_str}]")

    seq_str = ", ".join(f"{s['key']}={s['last_value']}" for s in bundle.tables["sequences"])
    lines += [
        "",
        "3. Đối soát tài chính & số lượng:",
        f"   - Tổng thành tiền (SL × Đơn giá)    : {rec.get('amount_after', 0):,.2f} (lệch: {rec.get('amount_diff', 0)})",
        f"   - Tổng lượng thực nhận (receipts)   : {rec.get('receipt_qty_after', 0):,.3f} (lệch: {rec.get('receipt_qty_diff', 0)})",
        f"   - Bộ đếm số phiếu (sequences)       : {seq_str}",
        "=" * 72,
    ]
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 5. Xuất file giao dịch SQL (--sql-out)
# ---------------------------------------------------------------------------


def sql_literal(v: Any) -> str:
    """Chuyển giá trị Python sang SQL literal chuẩn PostgreSQL (standard_conforming_strings = on)."""
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        d = format(Decimal(str(v)), "f")
        return d.rstrip("0").rstrip(".") if "." in d else d
    s = str(v).replace("\x00", "").replace("'", "''")
    return f"'{s}'"


def generate_sql(
    bundle: MigrationBundle,
    batch_size: int = 500,
    valid_auth_ids: set[str] | None = None,
) -> str:
    allowed_auth_ids = {u.lower() for u in valid_auth_ids} if valid_auth_ids else set()
    stmts: list[str] = [
        "-- File SQL Cutover được sinh tự động bởi scripts/migrate_to_supabase.py",
        "BEGIN;",
        "SET LOCAL TIME ZONE 'Asia/Ho_Chi_Minh';",
        "",
        (
            "TRUNCATE TABLE "
            + ", ".join(f"public.{t}" for t in reversed(SQL_ID_TABLES))
            + " RESTART IDENTITY CASCADE;"
        ),
        "",
    ]

    for tbl in SQL_ID_TABLES:
        cols = TABLE_COLUMNS[tbl]
        ts_cols = set(TIMESTAMPTZ_FIELDS.get(tbl, ()))
        rows = bundle.tables.get(tbl, [])
        if rows:
            quoted_cols = ", ".join(f'"{c}"' for c in cols)
            for i in range(0, len(rows), batch_size):
                chunk = rows[i : i + batch_size]
                values_lines: list[str] = []
                for r in chunk:
                    vals: list[str] = []
                    for c in cols:
                        val = r.get(c)
                        if c in ts_cols and val is not None:
                            val = ensure_vn_tz(val)
                        elif tbl == "receipts" and c == "created_by":
                            cb = parse_optional_uuid(val)
                            val = cb if (cb and cb in allowed_auth_ids) else None
                        vals.append(sql_literal(val))
                    values_lines.append("  (" + ", ".join(vals) + ")")
                stmts.append(
                    f"INSERT INTO public.{tbl} ({quoted_cols}) OVERRIDING SYSTEM VALUE VALUES\n"
                    + ",\n".join(values_lines)
                    + ";"
                )
            stmts.append("")

    stmts.append("-- Đồng bộ lại bộ đếm tự tăng (serial/identity sequence) cho mọi bảng có cột id")
    for tbl in SQL_ID_TABLES:
        stmts.append(
            f"SELECT setval(pg_get_serial_sequence('public.{tbl}', 'id'), "
            f"COALESCE((SELECT MAX(id) FROM public.{tbl}), 1), "
            f"(SELECT COUNT(*) > 0 FROM public.{tbl}));"
        )

    stmts.extend(["", "COMMIT;", ""])
    return "\n".join(stmts)


# ---------------------------------------------------------------------------
# 6. Nạp trực tiếp qua Supabase Auth Admin API & PostgREST (--execute)
# ---------------------------------------------------------------------------


def _http_json(
    method: str,
    url: str,
    headers: dict[str, str],
    payload: Any = None,
    timeout: int = 60,
) -> tuple[int, dict[str, str], Any]:
    body = None if payload is None else json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            raw_bytes = resp.read()
            resp_headers = {k.lower(): v for k, v in resp.headers.items()}
            data = json.loads(raw_bytes.decode("utf-8")) if raw_bytes else None
            return resp.status, resp_headers, data
    except urllib.error.HTTPError as exc:
        err_body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code} khi gọi {method} {url}: {err_body}") from exc


def _parse_content_range_total(content_range: str | None) -> int:
    if not content_range or "/" not in content_range:
        raise ValueError(f"Server không trả về Content-Range hợp lệ: {content_range!r}")
    total_part = content_range.split("/", 1)[1].strip()
    return int(total_part)


def execute_supabase(
    bundle: MigrationBundle,
    supabase_url: str,
    service_role_key: str,
    initial_password: str | None = None,
    batch_size: int = 500,
) -> dict[str, Any]:
    base_url = supabase_url.rstrip("/")
    if not base_url or not service_role_key:
        raise ValueError("Thiếu SUPABASE_URL hoặc SUPABASE_SERVICE_ROLE_KEY để chạy --execute")

    generated_pw = False
    if not initial_password:
        initial_password = secrets.token_urlsafe(12)
        generated_pw = True
    elif len(initial_password) < 8:
        raise ValueError("INITIAL_USER_PASSWORD phải có ít nhất 8 ký tự")

    email_domain = (
        os.environ.get("SYNTHETIC_EMAIL_DOMAIN", "dncvt.local").strip().lstrip("@")
        or "dncvt.local"
    )

    common_headers = {
        "apikey": service_role_key,
        "Authorization": f"Bearer {service_role_key}",
        "Content-Type": "application/json; charset=utf-8",
    }

    # 6.1 Tạo tài khoản Supabase Auth & nạp public.profiles
    _, _, existing_resp = _http_json(
        "GET",
        f"{base_url}/auth/v1/admin/users?page=1&per_page=1000",
        common_headers,
    )
    existing_users = existing_resp.get("users", []) if isinstance(existing_resp, dict) else []
    auth_by_email = {
        str(u.get("email", "")).lower(): str(u["id"]).lower()
        for u in existing_users
        if u.get("id")
    }
    valid_auth_ids: set[str] = {
        str(u["id"]).lower() for u in existing_users if u.get("id")
    }

    profiles_payload: list[dict[str, Any]] = []
    for u in bundle.tables["users"]:
        username = u["username"]
        email = f"{username.lower()}@{email_domain}"
        uid = auth_by_email.get(email)
        if not uid:
            auth_req_body: dict[str, Any] = {
                "email": email,
                "password": initial_password,
                "email_confirm": True,
                "user_metadata": {
                    "username": username,
                    "full_name": u["full_name"],
                    "role": u["role"],
                },
            }
            raw_uid = u.get("id")
            if isinstance(raw_uid, str) and UUID_RE.fullmatch(raw_uid) and raw_uid.lower() not in valid_auth_ids:
                auth_req_body["id"] = raw_uid.lower()
            if not u["active"]:
                auth_req_body["ban_duration"] = "876000h"
            _, _, created = _http_json(
                "POST",
                f"{base_url}/auth/v1/admin/users",
                common_headers,
                auth_req_body,
            )
            uid = str(created["id"]).lower()
            auth_by_email[email] = uid
            valid_auth_ids.add(uid)

        profiles_payload.append(
            {
                "id": uid,
                "username": username,
                "full_name": u["full_name"],
                "role": u["role"],
                "active": u["active"],
            }
        )

    delete_headers = {
        **common_headers,
        "Prefer": "return=minimal",
    }
    upsert_headers = {
        **common_headers,
        "Prefer": "resolution=merge-duplicates,return=minimal",
    }

    # 6.2 Xóa sạch dữ liệu seed/test cũ theo thứ tự ngược khóa ngoại trước khi nạp chuẩn
    for tbl in reversed(SQL_ID_TABLES):
        _http_json(
            "DELETE",
            f"{base_url}/rest/v1/{tbl}?id=gte.0",
            delete_headers,
        )

    if profiles_payload:
        keep_ids = ",".join(p["id"] for p in profiles_payload)
        _http_json(
            "DELETE",
            f"{base_url}/rest/v1/profiles?id=not.in.({keep_ids})",
            delete_headers,
        )
        _http_json(
            "POST",
            f"{base_url}/rest/v1/profiles?on_conflict=id",
            upsert_headers,
            profiles_payload,
        )
    else:
        _http_json(
            "DELETE",
            f"{base_url}/rest/v1/profiles?id=not.is.null",
            delete_headers,
        )

    # 6.3 Batch-upsert 500 dòng/lần cho toàn bộ 11 bảng dữ liệu
    for tbl in SQL_ID_TABLES:
        cols = TABLE_COLUMNS[tbl]
        ts_cols = set(TIMESTAMPTZ_FIELDS.get(tbl, ()))
        rows = bundle.tables.get(tbl, [])
        if not rows:
            continue
        for i in range(0, len(rows), batch_size):
            chunk = rows[i : i + batch_size]
            payload_chunk: list[dict[str, Any]] = []
            for r in chunk:
                item = {c: r.get(c) for c in cols}
                for c in ts_cols:
                    if item.get(c) is not None:
                        item[c] = ensure_vn_tz(item[c])
                if tbl == "requests" and isinstance(item["header_json"], str):
                    item["header_json"] = json.loads(item["header_json"])
                if tbl == "receipts":
                    cb = parse_optional_uuid(item.get("created_by"))
                    item["created_by"] = cb if (cb and cb in valid_auth_ids) else None
                payload_chunk.append(item)
            _http_json(
                "POST",
                f"{base_url}/rest/v1/{tbl}?on_conflict=id",
                upsert_headers,
                payload_chunk,
            )

    # 6.4 Đối soát số dòng thực tế trên server qua PostgREST
    count_headers = {
        **common_headers,
        "Prefer": "count=exact",
        "Range": "0-0",
    }
    server_counts: dict[str, int] = {}
    check_targets = [("profiles", len(profiles_payload))] + [
        (tbl, len(bundle.tables.get(tbl, []))) for tbl in SQL_ID_TABLES
    ]
    for tbl, expected in check_targets:
        _, resp_hdrs, _ = _http_json(
            "GET",
            f"{base_url}/rest/v1/{tbl}?select=id&limit=1",
            count_headers,
        )
        actual = _parse_content_range_total(resp_hdrs.get("content-range"))
        server_counts[tbl] = actual
        if actual != expected:
            raise ValueError(
                f"Đối soát trên server thất bại ở bảng {tbl}: server={actual}, kỳ vọng={expected}"
            )

    return {
        "server_counts": server_counts,
        "generated_password": initial_password if generated_pw else None,
        "users_provisioned": len(profiles_payload),
    }


# ---------------------------------------------------------------------------
# 7. Giao diện dòng lệnh (CLI)
# ---------------------------------------------------------------------------


def main(argv: list[str] | None = None) -> MigrationBundle:
    if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
        sys.stdout.reconfigure(encoding="utf-8")

    parser = argparse.ArgumentParser(
        description="Di trú & đối soát dữ liệu Đề nghị cấp vật tư sang Supabase (PostgreSQL)"
    )
    parser.add_argument("source", type=Path, help="Đường dẫn tới dev/data/db.json hoặc file .xlsx xuất từ admin.html")
    parser.add_argument("--dry-run", action="store_true", help="Chỉ chạy chuẩn hóa & in báo cáo đối soát (mặc định)")
    parser.add_argument("--sql-out", type=Path, default=None, help="Đường dẫn xuất file giao dịch SQL (.sql)")
    parser.add_argument("--execute", action="store_true", help="Nạp trực tiếp lên Supabase qua Auth Admin API & PostgREST")
    parser.add_argument("--batch-size", type=int, default=500, help="Số dòng mỗi lô upsert/insert (mặc định: 500)")

    args = parser.parse_args(argv)

    raw = load_source(args.source)
    bundle = transform_data(raw)
    reconcile(raw, bundle)

    print(format_report(bundle, source_name=args.source.name))

    if args.sql_out is not None:
        args.sql_out.parent.mkdir(parents=True, exist_ok=True)
        sql_content = generate_sql(bundle, batch_size=args.batch_size)
        args.sql_out.write_text(sql_content, encoding="utf-8")
        print(f"\n[SQL] Đã xuất file giao dịch SQL an toàn tại: {args.sql_out}")

    if args.execute:
        sb_url = os.environ.get("SUPABASE_URL", "").strip()
        sb_key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "").strip()
        init_pw = os.environ.get("INITIAL_USER_PASSWORD", "").strip() or None
        exec_res = execute_supabase(
            bundle,
            supabase_url=sb_url,
            service_role_key=sb_key,
            initial_password=init_pw,
            batch_size=args.batch_size,
        )
        if exec_res["generated_password"]:
            print(
                f"\n[AUTH] Mật khẩu khởi tạo tạm thời cho {exec_res['users_provisioned']} tài khoản: "
                f"{exec_res['generated_password']} (yêu cầu đổi mật khẩu ngay sau lần đăng nhập đầu tiên)"
            )
        print(f"\n[EXECUTE] Đã nạp & đối soát khớp 100% trên Supabase: {exec_res['server_counts']}")

    return bundle


if __name__ == "__main__":
    main()
