from __future__ import annotations

import copy
import json
import subprocess
import sys
import threading
import unicodedata
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import openpyxl
import pytest

WORKTREE_ROOT = Path(__file__).resolve().parent.parent
MAIN_REPO_DB_JSON = WORKTREE_ROOT.parent.parent / "dev" / "data" / "db.json"

sys.path.insert(0, str(WORKTREE_ROOT / "scripts"))
import migrate_to_supabase as mig  # noqa: E402


# ---------------------------------------------------------------------------
# Fixtures dữ liệu tổng hợp chứa đủ 8 dị thường lịch sử
# ---------------------------------------------------------------------------


def _make_synthetic_db_dict() -> dict:
    """Tạo cấu trúc db.json tổng hợp chứa đủ 8 dị thường dữ liệu lịch sử."""
    nfd_unit = unicodedata.normalize("NFD", "P.Cải tiến")
    nfd_name = unicodedata.normalize("NFD", "Ống thép đúc đặc biệt")
    header_ms = json.dumps(
        {
            "unit": "P.Cải tiến",
            "project_code": "Ms287",
            "project_name": "Cải tạo hệ thống đúc",
            "delivery_location": "Kho 1",
            "requester": "Hà",
            "request_date": "2026-10-20",
            "need_date": "2026-10-25",
            "note": "Giữ nguyên Ms287 trong header_json!",
        },
        ensure_ascii=False,
    )
    header_ls = json.dumps(
        {
            "unit": "Phòng R&D",
            "project_code": "MS027",
            "project_name": "Hệ thống Lò NL",
            "delivery_location": "",
            "requester": "Nam",
            "request_date": None,
            "need_date": "2021-12-23",
            "note": "",
        },
        ensure_ascii=False,
    )

    return {
        "tabs": {
            "lines": {
                "headers": [
                    "id", "request_code", "unit", "material_group", "name", "uom",
                    "qty_requested", "qty_on_hand", "need_date", "project_code",
                    "project_name", "purpose", "equipment_name", "delivery_location",
                    "requester", "request_date", "supplier", "unit_price", "note",
                    "legacy_flag", "created_by", "created_at", "updated_at",
                ],
                "rows": [
                    [
                        "10", "LS-0001", nfd_unit, "C", nfd_name, "cây",
                        "5", "1", "2021-12-23", "Ms027",
                        "Hệ thống Lò NL", "Sửa lò", "", "PX Cơ điện",
                        "Nam", "2021-12-20", "KHVT", "120000", "Ghi chú 'nháy đơn' \\ gạch chéo",
                        "TRUE", "import", "2021-12-20T00:00:00", "2022-01-15T08:30:00",
                    ],
                    [
                        "25", "YC2026-150", "P.Cải tiến", "E", "=SUM(A1:A5); '; DROP TABLE public.lines; --", "3",
                        "2.5", "", "2026-10-25", " ms287 ",
                        "Cải tạo hệ thống đúc", "Thay mới", "Máy đúc", "Kho 1",
                        "Hà", "2026-10-20", "CK-DNCO1", "40000.5", "",
                        "", "ha", "2026-10-20T09:00:00", "2026-10-21T14:22:11",
                    ],
                    [
                        "99", "", "Ban dự án", "C", "Bulong M16", "bộ",
                        "10", "0", "2026-10-30", "DA-999",
                        "Dự án mở rộng", "", "", "",
                        "Nam", "", "", "", "",
                        "", "nam", "2026-10-25T10:00:00", "2026-10-25T10:00:00",
                    ],
                ],
            },
            "requests": {
                "headers": [
                    "id", "code", "year", "seq", "issued_at", "issued_by",
                    "header_json", "line_count", "legacy_flag",
                ],
                "rows": [
                    ["1", "LS-0001", "2021", "1", "", "Nam", header_ls, "1", "TRUE"],
                    ["2", "YC2026-150", "2026", "150", "2026-10-20T09:05:00", "Hà", header_ms, "1", ""],
                    ["3", "YC2025-400", "2025", "400", "2025-12-31T16:00:00", "Nam", header_ls, "0", ""],
                ],
            },
            "receipts": {
                "headers": [
                    "id", "line_id", "received_date", "qty", "received_by", "note", "created_at",
                ],
                "rows": [
                    ["1", "10", "2021-12-22", "2.1114", "Nam", "Đợt 1", "2021-12-22T10:00:00"],
                    ["2", "10", "2021-12-25", "1.2224", "Nam", "Đợt 2", "2021-12-25T15:00:00"],
                    ["3", "25", "", "1.5", "", "import; thiếu ngày nhận", "2026-10-22T00:00:00"],
                ],
            },
            "suppliers": {
                "headers": ["id", "code", "name", "address", "phone", "email", "tax_code", "field"],
                "rows": [
                    ["1", "CK-DNCO1", "Công ty Đức Nga", "Nam Định", "0947904546", "", "", "Cơ khí"],
                ],
            },
            "project_items": {
                "headers": ["id", "project_code", "project_name", "item_code", "item_name", "owner", "status"],
                "rows": [
                    # Cặp trùng 1: id=27 thiếu owner/status, id=369 có owner/status & tên đầy đủ hơn -> giữ id=27, gộp cột đầy đủ
                    ["369", " ms027 ", "Dự án Lò Năng Lượng", "", "Hệ thống Lò NL khu A", "Nguyễn Văn A", "Đang thực hiện"],
                    ["27", "MS027", "", "", "Hệ thống Lò NL", "", ""],
                    # Cặp trùng 2: viết thường Ms287
                    ["287", "Ms287", "", "", "Cải tạo hệ thống đúc chân không", "Trần B", ""],
                    ["373", "MS287", "", "", "Cải tạo hệ thống đúc chân không", "", "Hoàn thành"],
                    # Mã không phải MS\d+ giữ nguyên chữ hoa/thường vốn có
                    ["500", "DA-999", "Dự án mở rộng", "HM-01", "Hạng mục 1", "", ""],
                ],
            },
            "material_groups": {
                "headers": ["id", "name", "description"],
                "rows": [
                    ["1", "C", "Thiết bị, chế tạo"],
                    ["5", "E", "Văn phòng phẩm"],
                ],
            },
            "units": {
                "headers": ["id", "name"],
                "rows": [
                    ["1", "Ban dự án"],
                    ["2", "P.Cải tiến"],
                ],
            },
            "uoms": {
                "headers": ["id", "name"],
                "rows": [
                    ["1", "3"],
                    ["2", "cây"],
                    ["3", "bộ"],
                ],
            },
            # Cố tình không có tab 'signers' để kiểm tra tự động seed 8 dòng DEFAULT_SIGNERS
            "users": {
                "headers": ["id", "username", "full_name", "role", "password_hash", "salt", "active", "created_at"],
                "rows": [
                    ["1", "admin", "Quản trị", "admin", "hash1", "salt1", "TRUE", "2026-10-23T09:13:54"],
                    ["2", "nam", "Nguyễn Văn Nam", "user", "hash2", "salt2", "TRUE", "2026-10-23T09:14:03"],
                ],
            },
            "audit_log": {
                "headers": ["id", "at", "user", "table", "row_id", "field", "old", "new"],
                "rows": [],
            },
            "sequences": {
                "headers": ["id", "key", "last_value"],
                "rows": [
                    ["1", "YC2025", "391"],  # nhỏ hơn seq=400 trong requests -> phải nâng lên 400
                    ["2", "YC2026", "100"],  # nhỏ hơn INITIAL_SEQ[2026]=294 -> phải nâng lên 294
                ],
            },
        },
        "props": {},
    }


@pytest.fixture
def synthetic_json_path(tmp_path: Path) -> Path:
    p = tmp_path / "synthetic_db.json"
    p.write_text(json.dumps(_make_synthetic_db_dict(), ensure_ascii=False), encoding="utf-8")
    return p


def _write_admin_export_xlsx(db_dict: dict, out_path: Path, include_seeded_signers: bool = False) -> Path:
    """Mô phỏng chính xác file ncvt-toan-bo-YYYY-MM-DD.xlsx xuất từ admin.html (12 sheets)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    schema_order = [
        "lines", "requests", "receipts", "suppliers", "project_items",
        "material_groups", "signers", "units", "uoms", "users", "audit_log", "sequences",
    ]
    bool_fields = {
        ("lines", "legacy_flag"),
        ("requests", "legacy_flag"),
        ("users", "active"),
    }
    int_fields = {
        "id", "year", "seq", "line_count", "line_id", "last_value",
    }
    num_fields = {"qty_requested", "qty_on_hand", "unit_price", "qty"}

    for name in schema_order:
        ws = wb.create_sheet(name)
        tab = db_dict["tabs"].get(name)
        if name == "signers" and (not tab or not tab["rows"]):
            if include_seeded_signers:
                ws.append(["id", "unit_contains", "position", "title", "name"])
                for idx, s in enumerate(mig.DEFAULT_SIGNERS, 1):
                    ws.append([idx, s["unit_contains"] or "", s["position"], s["title"], s["name"] or ""])
            else:
                # Bảng trống trên admin.html chỉ ghi header ['id']
                ws.append(["id"])
            continue

        if not tab or not tab["rows"]:
            ws.append(["id"])
            continue

        headers = list(tab["headers"])
        if name == "users":
            # apiExportAll_() gọi apiUsersList_() -> bỏ password_hash, salt, created_at
            keep_cols = ["id", "username", "full_name", "role", "active"]
            col_idx = [headers.index(c) for c in keep_cols]
            ws.append(keep_cols)
            for r in tab["rows"]:
                out_row = []
                for c, idx in zip(keep_cols, col_idx):
                    v = r[idx] if idx < len(r) else ""
                    if c == "id":
                        out_row.append(int(v))
                    elif c == "active":
                        out_row.append(v in (True, "TRUE", "true", "1", 1))
                    else:
                        out_row.append(v)
                ws.append(out_row)
            continue

        ws.append(headers)
        for r in tab["rows"]:
            out_row = []
            for c, v in zip(headers, r):
                if (name, c) in bool_fields:
                    out_row.append(v in (True, "TRUE", "true", "1", 1))
                elif v == "" or v is None:
                    out_row.append("")
                elif c in int_fields:
                    out_row.append(int(v))
                elif c in num_fields:
                    fv = float(v)
                    out_row.append(int(fv) if fv.is_integer() else fv)
                else:
                    out_row.append(v)
            ws.append(out_row)

    wb.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# 1. Test chuẩn hóa 8 dị thường lịch sử trên bộ dữ liệu tổng hợp & db.json thật
# ---------------------------------------------------------------------------


def test_synthetic_8_anomalies_and_reconcile(synthetic_json_path: Path):
    raw = mig.load_source(synthetic_json_path)
    bundle = mig.transform_data(raw)
    summary = mig.reconcile(raw, bundle)

    # 1) project_items: 5 dòng trước -> gộp 2 cặp trùng còn 3 dòng; giữ id nhỏ hơn & gộp cột đầy đủ nhất; Ms -> MS
    pi_by_id = {r["id"]: r for r in bundle.tables["project_items"]}
    assert set(pi_by_id.keys()) == {27, 287, 500}
    assert len(bundle.merged_project_items) == 2

    assert pi_by_id[27]["project_code"] == "MS027"
    assert pi_by_id[27]["project_name"] == "Dự án Lò Năng Lượng"
    assert pi_by_id[27]["item_name"] == "Hệ thống Lò NL khu A"
    assert pi_by_id[27]["owner"] == "Nguyễn Văn A"
    assert pi_by_id[27]["status"] == "Đang thực hiện"

    assert pi_by_id[287]["project_code"] == "MS287"
    assert pi_by_id[287]["owner"] == "Trần B"
    assert pi_by_id[287]["status"] == "Hoàn thành"

    assert pi_by_id[500]["project_code"] == "DA-999"

    # 2) lines: giữ nguyên id (10, 25, 99), chuẩn hóa Ms->MS, chuẩn hóa NFC, giữ nguyên created_at/updated_at
    lines_by_id = {r["id"]: r for r in bundle.tables["lines"]}
    assert set(lines_by_id.keys()) == {10, 25, 99}
    assert lines_by_id[10]["project_code"] == "MS027"
    assert lines_by_id[25]["project_code"] == "MS287"
    assert lines_by_id[99]["project_code"] == "DA-999"

    # Kiểm tra NFC
    assert lines_by_id[10]["unit"] == "P.Cải tiến"
    assert unicodedata.is_normalized("NFC", lines_by_id[10]["unit"])
    assert lines_by_id[10]["name"] == "Ống thép đúc đặc biệt"
    assert unicodedata.is_normalized("NFC", lines_by_id[10]["name"])

    # Giữ nguyên tuyệt đối created_at & updated_at lịch sử
    assert lines_by_id[10]["created_at"] == "2021-12-20T00:00:00"
    assert lines_by_id[10]["updated_at"] == "2022-01-15T08:30:00"
    assert lines_by_id[25]["updated_at"] == "2026-10-21T14:22:11"

    # Tính sẵn qty_received (round 3 chữ số) và last_received_date
    assert lines_by_id[10]["qty_received"] == 3.334  # round(2.1114 + 1.2224, 3)
    assert lines_by_id[10]["last_received_date"] == "2021-12-25"
    assert lines_by_id[25]["qty_received"] == 1.5
    assert lines_by_id[25]["last_received_date"] is None  # receipt không có ngày nhận
    assert lines_by_id[99]["qty_received"] == 0
    assert lines_by_id[99]["last_received_date"] is None

    # 3) requests: IMMUTABLE header_json (giữ nguyên cả "Ms287" bên trong!) & cho phép issued_at = None
    reqs_by_code = {r["code"]: r for r in bundle.tables["requests"]}
    raw_reqs_by_code = {r["code"]: r for r in raw["requests"]}
    assert reqs_by_code["LS-0001"]["issued_at"] is None
    assert reqs_by_code["YC2026-150"]["header_json"] == raw_reqs_by_code["YC2026-150"]["header_json"]
    assert json.loads(reqs_by_code["YC2026-150"]["header_json"])["project_code"] == "Ms287"

    # 4) receipts: giữ nguyên received_date=None, created_by=None
    recs_by_id = {r["id"]: r for r in bundle.tables["receipts"]}
    assert recs_by_id[3]["received_date"] is None
    assert all(r["created_by"] is None for r in bundle.tables["receipts"])

    # 5) signers: tự động seed 8 dòng DEFAULT_SIGNERS
    assert len(bundle.tables["signers"]) == 8
    assert bundle.tables["signers"][0]["title"] == "Người lập phiếu"
    assert bundle.tables["signers"][4]["unit_contains"] == "Cải tiến"

    # 6) sequences: YC2025 nâng từ 391 -> 400 (theo MAX(seq)), YC2026 nâng từ 100 -> 294 (theo INITIAL_SEQ[2026])
    seqs = {s["key"]: s["last_value"] for s in bundle.tables["sequences"]}
    assert seqs["YC2025"] == 400
    assert seqs["YC2026"] == 294

    # 7) Báo cáo đối soát tiếng Việt
    report = mig.format_report(bundle, source_name="synthetic_db.json")
    assert "PASS" in report
    assert "MS027" in report and "MS287" in report
    assert summary["amount_diff"] == 0
    assert summary["receipt_qty_diff"] == 0


def test_invalid_header_json_raises_error(synthetic_json_path: Path):
    raw = mig.load_source(synthetic_json_path)
    raw["requests"][0]["header_json"] = "{invalid json"
    with pytest.raises(ValueError, match="header_json"):
        mig.transform_data(raw)


def test_existing_signers_not_overwritten(synthetic_json_path: Path):
    raw = mig.load_source(synthetic_json_path)
    raw["signers"] = [
        {"id": "1", "unit_contains": "", "position": "1", "title": "Người duyệt riêng", "name": "Nguyễn X"}
    ]
    bundle = mig.transform_data(raw)
    assert len(bundle.tables["signers"]) == 1
    assert bundle.tables["signers"][0]["name"] == "Nguyễn X"


@pytest.mark.skipif(not MAIN_REPO_DB_JSON.exists(), reason="Không tìm thấy dev/data/db.json thật")
def test_real_db_json_full_migration_and_reconciliation():
    raw = mig.load_source(MAIN_REPO_DB_JSON)
    bundle = mig.transform_data(raw)
    summary = mig.reconcile(raw, bundle)

    # Kiểm tra số dòng thực tế của db.json
    assert len(bundle.tables["lines"]) == 7945
    assert len(bundle.tables["requests"]) == 2011
    assert len(bundle.tables["receipts"]) == 4654
    assert len(bundle.tables["suppliers"]) == 17
    assert len(raw["project_items"]) == 378
    assert len(bundle.tables["project_items"]) == 368
    assert len(bundle.merged_project_items) == 10
    assert len(bundle.tables["signers"]) == 8

    # Kiểm tra đúng 10 cặp trùng project_items đã gộp, giữ id nhỏ hơn
    merged_codes = {m["project_code"] for m in bundle.merged_project_items}
    assert merged_codes == {
        "MS027", "MS043", "MS222", "MS250", "MS287",
        "MS291", "MS315", "MS349", "MS350", "MS359",
    }
    kept_ids = {m["kept_id"] for m in bundle.merged_project_items}
    assert kept_ids == {27, 43, 222, 250, 287, 291, 315, 349, 350, 359}

    # Không còn bất kỳ mã Ms viết thường nào trong project_items và lines
    for row in bundle.tables["project_items"]:
        pc = row["project_code"] or ""
        assert not (pc.lower().startswith("ms") and pc[2:].isdigit() and pc != pc.upper())
    for row in bundle.tables["lines"]:
        pc = row["project_code"] or ""
        assert not (pc.lower().startswith("ms") and pc[2:].isdigit() and pc != pc.upper())

    # Giữ nguyên 100% requests.header_json (bao gồm YC2026-292 có "Ms367") và 39 phiếu trống issued_at
    raw_headers = {int(r["id"]): r["header_json"] for r in raw["requests"]}
    for req in bundle.tables["requests"]:
        assert req["header_json"] == raw_headers[req["id"]]
    null_issued = [r for r in bundle.tables["requests"] if r["issued_at"] is None]
    assert len(null_issued) == 39

    # Giữ nguyên 100% lines.created_at & updated_at
    raw_ts = {int(r["id"]): (r["created_at"], r["updated_at"]) for r in raw["lines"]}
    for line in bundle.tables["lines"]:
        assert (line["created_at"], line["updated_at"]) == raw_ts[line["id"]]

    # Đối soát thành tiền và lượng nhận lệch = 0
    assert summary["amount_diff"] == 0
    assert summary["receipt_qty_diff"] == 0
    assert round(sum(r["qty"] for r in bundle.tables["receipts"]), 3) == 232080.388

    # YC2026 >= 294
    seqs = {s["key"]: s["last_value"] for s in bundle.tables["sequences"]}
    assert seqs["YC2026"] >= 294

    # Sinh SQL từ toàn bộ 15.000+ dòng dữ liệu thật và nạp thử vào PostgreSQL (PGlite)
    sql_text = mig.generate_sql(bundle, batch_size=500)
    assert sql_text.lstrip().startswith("--")
    assert "\nBEGIN;\n" in sql_text and sql_text.rstrip().endswith("COMMIT;")


@pytest.mark.skipif(not MAIN_REPO_DB_JSON.exists(), reason="Không tìm thấy dev/data/db.json thật")
def test_real_db_json_sql_executes_cleanly_in_pglite(tmp_path: Path):
    pglite_dir = WORKTREE_ROOT / "node_modules" / "@electric-sql" / "pglite"
    if not pglite_dir.exists():
        pytest.skip("Không có @electric-sql/pglite")

    sql_path = tmp_path / "real_db_cutover.sql"
    mig.main([str(MAIN_REPO_DB_JSON), "--sql-out", str(sql_path)])

    js_runner = tmp_path / "verify_real_sql.js"
    js_runner.write_text(
        f"""
        const fs = require('fs');
        const {{ PGlite }} = require({json.dumps(str(pglite_dir))});
        (async () => {{
          const db = new PGlite();
          await db.exec(`
            CREATE SCHEMA IF NOT EXISTS private;
            CREATE OR REPLACE FUNCTION private.fmt_ts(p_ts timestamptz) RETURNS text LANGUAGE sql STABLE AS $$
              SELECT CASE WHEN p_ts IS NULL THEN NULL ELSE to_char(p_ts AT TIME ZONE 'Asia/Ho_Chi_Minh', 'YYYY-MM-DD"T"HH24:MI:SS') END;
            $$;
            CREATE TABLE public.suppliers (id bigint generated by default as identity primary key, code text, name text, address text, phone text, email text, tax_code text, field text);
            CREATE TABLE public.project_items (id bigint generated by default as identity primary key, project_code text unique, project_name text, item_code text, item_name text, owner text, status text);
            CREATE TABLE public.material_groups (id bigint generated by default as identity primary key, name text unique, description text);
            CREATE TABLE public.units (id bigint generated by default as identity primary key, name text unique);
            CREATE TABLE public.uoms (id bigint generated by default as identity primary key, name text unique);
            CREATE TABLE public.signers (id bigint generated by default as identity primary key, unit_contains text, position text, title text, name text);
            CREATE TABLE public.requests (id bigint generated by default as identity primary key, code text unique not null, year int not null, seq int not null, issued_at timestamptz, issued_by text, header_json jsonb not null, line_count int not null default 0, legacy_flag boolean not null default false);
            CREATE TABLE public.lines (id bigint generated by default as identity primary key, request_code text references public.requests(code), unit text, material_group text, name text, uom text, qty_requested numeric, qty_on_hand numeric, need_date date, project_code text, project_name text, purpose text, equipment_name text, delivery_location text, requester text, request_date date, supplier text, unit_price numeric, note text, legacy_flag boolean not null default false, created_by text, created_at timestamptz, updated_at timestamptz, qty_received numeric not null default 0, last_received_date date);
            CREATE TABLE public.receipts (id bigint generated by default as identity primary key, line_id bigint not null references public.lines(id), received_date date, qty numeric not null, received_by text, note text, created_at timestamptz, created_by uuid);
            CREATE TABLE public.audit_log (id bigint generated by default as identity primary key, at timestamptz, "user" text, "table" text, row_id text, field text, "old" text, "new" text);
            CREATE TABLE public.sequences (id bigint generated by default as identity primary key, key text unique not null, last_value bigint not null);
            SET TIME ZONE 'UTC';
          `);
          await db.exec(fs.readFileSync({json.dumps(str(sql_path))}, 'utf8'));
          const counts = {{}};
          for (const t of ['lines', 'requests', 'receipts', 'project_items', 'signers', 'sequences']) {{
            counts[t] = Number((await db.query(`SELECT COUNT(*) AS c FROM public.${{t}}`)).rows[0].c);
          }}
          const recSum = Number((await db.query('SELECT ROUND(SUM(qty), 3)::float8 AS s FROM public.receipts')).rows[0].s);
          const nextLine = Number((await db.query("SELECT nextval(pg_get_serial_sequence('public.lines', 'id')) AS n")).rows[0].n);
          const line1Upd = (await db.query("SELECT private.fmt_ts(updated_at) AS u FROM public.lines WHERE id = 1")).rows[0].u;
          await db.close();
          console.log(JSON.stringify({{ counts, recSum, nextLine, line1Upd }}));
        }})().catch(e => {{ console.error(e); process.exit(1); }});
        """,
        encoding="utf-8",
    )
    proc = subprocess.run(["node", str(js_runner)], capture_output=True, text=True, check=True, encoding="utf-8")
    res = json.loads(proc.stdout.strip().splitlines()[-1])
    assert res["counts"] == {
        "lines": 7945,
        "requests": 2011,
        "receipts": 4654,
        "project_items": 368,
        "signers": 8,
        "sequences": 6,
    }
    assert res["recSum"] == 232080.388
    assert res["nextLine"] == 7946
    assert res["line1Upd"] == "2026-10-26T09:30:34"


# ---------------------------------------------------------------------------
# 2. Test đọc từ file .xlsx 12-sheet đồng nhất 100% với đọc từ db.json
# ---------------------------------------------------------------------------


def test_xlsx_12_sheets_matches_json_identically(tmp_path: Path):
    db_dict = _make_synthetic_db_dict()
    json_path = tmp_path / "db.json"
    json_path.write_text(json.dumps(db_dict, ensure_ascii=False), encoding="utf-8")

    xlsx_empty_signers = _write_admin_export_xlsx(db_dict, tmp_path / "ncvt-toan-bo-2026-10-26.xlsx", include_seeded_signers=False)
    xlsx_with_signers = _write_admin_export_xlsx(db_dict, tmp_path / "ncvt-seeded.xlsx", include_seeded_signers=True)

    b_json = mig.transform_data(mig.load_source(json_path))
    b_xlsx1 = mig.transform_data(mig.load_source(xlsx_empty_signers))
    b_xlsx2 = mig.transform_data(mig.load_source(xlsx_with_signers))

    mig.reconcile(mig.load_source(xlsx_empty_signers), b_xlsx1)
    mig.reconcile(mig.load_source(xlsx_with_signers), b_xlsx2)

    compare_tables = [
        "lines", "requests", "receipts", "suppliers", "project_items",
        "material_groups", "units", "uoms", "signers", "audit_log", "sequences",
    ]
    for tbl in compare_tables:
        assert b_xlsx1.tables[tbl] == b_json.tables[tbl], f"Lệch bảng {tbl} (xlsx_empty_signers)"
        assert b_xlsx2.tables[tbl] == b_json.tables[tbl], f"Lệch bảng {tbl} (xlsx_with_signers)"

    # Bảng users từ admin.html không chứa password_hash/salt nhưng khớp toàn bộ thông tin tài khoản
    for u_xlsx, u_json in zip(b_xlsx1.tables["users"], b_json.tables["users"]):
        for k in ("id", "username", "full_name", "role", "active"):
            assert u_xlsx[k] == u_json[k]


# ---------------------------------------------------------------------------
# 3. Test --sql-out sinh giao dịch SQL hợp lệ, chống SQL injection & chạy thật trên PGlite
# ---------------------------------------------------------------------------


def test_sql_out_valid_and_executes_in_pglite(synthetic_json_path: Path, tmp_path: Path):
    sql_path = tmp_path / "cutover.sql"
    bundle = mig.main([str(synthetic_json_path), "--sql-out", str(sql_path)])
    assert bundle is not None
    assert sql_path.exists()

    sql_text = sql_path.read_text(encoding="utf-8")
    assert sql_text.lstrip().startswith("BEGIN;") or "\nBEGIN;\n" in sql_text
    assert sql_text.rstrip().endswith("COMMIT;")
    assert "OVERRIDING SYSTEM VALUE" in sql_text

    # Có setval đầy đủ cho toàn bộ 11 bảng có cột id
    for tbl in mig.SQL_ID_TABLES:
        expected_setval = (
            f"SELECT setval(pg_get_serial_sequence('public.{tbl}', 'id'), "
            f"COALESCE((SELECT MAX(id) FROM public.{tbl}), 1), "
            f"(SELECT COUNT(*) > 0 FROM public.{tbl}));"
        )
        assert expected_setval in sql_text, f"Thiếu setval chuẩn cho bảng {tbl}"

    # Chạy thật file SQL vừa xuất trên PostgreSQL (PGlite) để chứng minh cú pháp & chống SQL injection
    pglite_dir = WORKTREE_ROOT / "node_modules" / "@electric-sql" / "pglite"
    if pglite_dir.exists():
        js_runner = tmp_path / "verify_sql.js"
        js_runner.write_text(
            f"""
            const fs = require('fs');
            const {{ PGlite }} = require({json.dumps(str(pglite_dir))});
            (async () => {{
              const db = new PGlite();
              await db.exec(`
                CREATE TABLE public.suppliers (id bigint generated by default as identity primary key, code text, name text, address text, phone text, email text, tax_code text, field text);
                CREATE TABLE public.project_items (id bigint generated by default as identity primary key, project_code text unique, project_name text, item_code text, item_name text, owner text, status text);
                CREATE TABLE public.material_groups (id bigint generated by default as identity primary key, name text unique, description text);
                CREATE TABLE public.units (id bigint generated by default as identity primary key, name text unique);
                CREATE TABLE public.uoms (id bigint generated by default as identity primary key, name text unique);
                CREATE TABLE public.signers (id bigint generated by default as identity primary key, unit_contains text, position text, title text, name text);
                CREATE TABLE public.requests (id bigint generated by default as identity primary key, code text unique not null, year int not null, seq int not null, issued_at timestamptz, issued_by text, header_json jsonb not null, line_count int not null default 0, legacy_flag boolean not null default false);
                CREATE TABLE public.lines (id bigint generated by default as identity primary key, request_code text references public.requests(code), unit text, material_group text, name text, uom text, qty_requested numeric, qty_on_hand numeric, need_date date, project_code text, project_name text, purpose text, equipment_name text, delivery_location text, requester text, request_date date, supplier text, unit_price numeric, note text, legacy_flag boolean not null default false, created_by text, created_at timestamptz, updated_at timestamptz, qty_received numeric not null default 0, last_received_date date);
                CREATE TABLE public.receipts (id bigint generated by default as identity primary key, line_id bigint not null references public.lines(id), received_date date, qty numeric not null, received_by text, note text, created_at timestamptz, created_by uuid);
                CREATE TABLE public.audit_log (id bigint generated by default as identity primary key, at timestamptz, "user" text, "table" text, row_id text, field text, "old" text, "new" text);
                CREATE TABLE public.sequences (id bigint generated by default as identity primary key, key text unique not null, last_value bigint not null);
              `);
              const sql = fs.readFileSync({json.dumps(str(sql_path))}, 'utf8');
              await db.exec(sql);
              // Chạy lần 2 để kiểm tra tính idempotent của file SQL
              await db.exec(sql);

              const l10 = (await db.query('SELECT note, qty_received::float8 AS qr FROM public.lines WHERE id = 10')).rows[0];
              const l25 = (await db.query('SELECT name, uom FROM public.lines WHERE id = 25')).rows[0];
              const req2 = (await db.query("SELECT header_json->>'project_code' AS pc FROM public.requests WHERE code = 'YC2026-150'")).rows[0];
              const nextLineId = (await db.query("SELECT nextval(pg_get_serial_sequence('public.lines', 'id')) AS nid")).rows[0];
              const nextAuditId = (await db.query("SELECT nextval(pg_get_serial_sequence('public.audit_log', 'id')) AS aid")).rows[0];
              await db.close();
              console.log(JSON.stringify({{ l10, l25, req2, nextLineId, nextAuditId }}));
            }})().catch(e => {{ console.error(e); process.exit(1); }});
            """,
            encoding="utf-8",
        )
        proc = subprocess.run(["node", str(js_runner)], capture_output=True, text=True, check=True, encoding="utf-8")
        out = json.loads(proc.stdout.strip().splitlines()[-1])
        assert out["l10"]["note"] == "Ghi chú 'nháy đơn' \\ gạch chéo"
        assert out["l10"]["qr"] == 3.334
        assert out["l25"]["name"] == "=SUM(A1:A5); '; DROP TABLE public.lines; --"
        assert out["l25"]["uom"] == "3"
        assert out["req2"]["pc"] == "Ms287"
        assert int(out["nextLineId"]["nid"]) == 100  # MAX(id)=99 -> nextval=100
        assert int(out["nextAuditId"]["aid"]) == 1   # bảng rỗng -> nextval=1


# ---------------------------------------------------------------------------
# 4. Test reconcile phát hiện sai lệch (cố tình làm lệch -> raise ValueError)
# ---------------------------------------------------------------------------


def test_reconcile_detects_tampering(synthetic_json_path: Path):
    raw = mig.load_source(synthetic_json_path)
    base_bundle = mig.transform_data(raw)

    # 4a. Lệch 1 đồng đơn giá / thành tiền
    b_amt = copy.deepcopy(base_bundle)
    b_amt.tables["lines"][0]["unit_price"] += 1
    with pytest.raises(ValueError, match="thành tiền|amount"):
        mig.reconcile(raw, b_amt)

    # 4b. Thiếu 1 dòng vật tư
    b_line = copy.deepcopy(base_bundle)
    b_line.tables["lines"].pop()
    with pytest.raises(ValueError, match="lines"):
        mig.reconcile(raw, b_line)

    # 4c. Đổi id gốc (STT) của dòng vật tư
    b_id = copy.deepcopy(base_bundle)
    b_id.tables["lines"][0]["id"] = 999999
    with pytest.raises(ValueError, match="id|lines"):
        mig.reconcile(raw, b_id)

    # 4d. Ghi đè updated_at lịch sử của lines
    b_ts = copy.deepcopy(base_bundle)
    b_ts.tables["lines"][0]["updated_at"] = "2099-01-01T00:00:00"
    with pytest.raises(ValueError, match="updated_at|created_at"):
        mig.reconcile(raw, b_ts)

    # 4e. Sửa 1 ký tự trong requests.header_json (vd. đổi Ms287 -> MS287 trong header_json)
    b_hdr = copy.deepcopy(base_bundle)
    b_hdr.tables["requests"][1]["header_json"] = b_hdr.tables["requests"][1]["header_json"].replace("Ms287", "MS287")
    with pytest.raises(ValueError, match="header_json"):
        mig.reconcile(raw, b_hdr)

    # 4f. Lệch 0.001 tổng lượng nhận receipts
    b_rec = copy.deepcopy(base_bundle)
    b_rec.tables["receipts"][0]["qty"] += 0.001
    with pytest.raises(ValueError, match="receipts|nhận"):
        mig.reconcile(raw, b_rec)

    # 4g. Sót mã trùng hoặc lệch tổng số dòng project_items
    b_pi = copy.deepcopy(base_bundle)
    b_pi.tables["project_items"].append(
        {"id": 999, "project_code": "MS027", "project_name": None, "item_code": None, "item_name": "Trùng", "owner": None, "status": None}
    )
    with pytest.raises(ValueError, match="project_items"):
        mig.reconcile(raw, b_pi)


# ---------------------------------------------------------------------------
# 5. Test chế độ --dry-run và --execute (Auth Admin API + PostgREST batch 500)
# ---------------------------------------------------------------------------


def test_dry_run_default_writes_nothing(synthetic_json_path: Path, capsys):
    mig.main([str(synthetic_json_path)])
    captured = capsys.readouterr().out
    assert "BÁO CÁO ĐỐI SOÁT" in captured
    assert "PASS" in captured


def test_execute_mode_batches_500_and_creates_auth_users(tmp_path: Path, monkeypatch):
    # Tạo dữ liệu có 505 dòng lines để kiểm tra chia lô 500 + 5 qua PostgREST
    db_dict = _make_synthetic_db_dict()
    base_row = db_dict["tabs"]["lines"]["rows"][2]
    db_dict["tabs"]["lines"]["rows"] = [
        [str(i)] + base_row[1:] for i in range(1, 506)
    ]
    db_dict["tabs"]["requests"]["rows"] = []
    db_dict["tabs"]["receipts"]["rows"] = []
    src_path = tmp_path / "bulk.json"
    src_path.write_text(json.dumps(db_dict, ensure_ascii=False), encoding="utf-8")

    state = {
        "auth_users": [],
        "batches": {},  # table -> list of batch lengths
        "rows": {},     # table -> list of rows
    }

    class FakeSupabaseHandler(BaseHTTPRequestHandler):
        def log_message(self, fmt, *args):
            pass

        def _send_json(self, status: int, payload, extra_headers: dict | None = None):
            body = json.dumps(payload).encode("utf-8")
            self.send_response(status)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            for k, v in (extra_headers or {}).items():
                self.send_header(k, v)
            self.end_headers()
            self.wfile.write(body)

        def do_GET(self):
            parsed = urlparse(self.path)
            if parsed.path == "/auth/v1/admin/users":
                self._send_json(200, {"users": state["auth_users"]})
                return
            if parsed.path.startswith("/rest/v1/"):
                tbl = parsed.path.split("/rest/v1/")[1]
                cnt = len(state["rows"].get(tbl, []))
                rng = f"0-0/{cnt}" if cnt > 0 else "*/0"
                self._send_json(200, [], {"Content-Range": rng})
                return
            self._send_json(404, {"error": "not found"})

        def do_POST(self):
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length).decode("utf-8")) if length else None
            parsed = urlparse(self.path)
            if parsed.path == "/auth/v1/admin/users":
                uid = f"00000000-0000-0000-0000-{len(state['auth_users']) + 1:012d}"
                u = {"id": uid, "email": payload["email"], "email_confirm": payload.get("email_confirm")}
                state["auth_users"].append(u)
                self._send_json(200, u)
                return
            if parsed.path.startswith("/rest/v1/"):
                tbl = parsed.path.split("/rest/v1/")[1]
                items = payload if isinstance(payload, list) else [payload]
                state["batches"].setdefault(tbl, []).append(len(items))
                state["rows"].setdefault(tbl, []).extend(items)
                self._send_json(201, [])
                return
            self._send_json(404, {"error": "not found"})

        def do_DELETE(self):
            parsed = urlparse(self.path)
            if parsed.path.startswith("/rest/v1/"):
                tbl = parsed.path.split("/rest/v1/")[1]
                state["rows"][tbl] = []
                self._send_json(200, [])
                return
            self._send_json(404, {"error": "not found"})

    server = ThreadingHTTPServer(("127.0.0.1", 0), FakeSupabaseHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        base_url = f"http://127.0.0.1:{server.server_port}"
        monkeypatch.setenv("SUPABASE_URL", base_url)
        monkeypatch.setenv("SUPABASE_SERVICE_ROLE_KEY", "test-service-role-key")
        monkeypatch.setenv("INITIAL_USER_PASSWORD", "MatKhauTam@2026")

        mig.main([str(src_path), "--execute"])

        assert [u["email"] for u in state["auth_users"]] == ["admin@dncvt.local", "nam@dncvt.local"]
        assert all(u["email_confirm"] is True for u in state["auth_users"])
        assert len(state["rows"]["profiles"]) == 2
        assert state["batches"]["lines"] == [500, 5]  # Chia đúng lô 500 dòng/lần
        assert len(state["rows"]["signers"]) == 8
    finally:
        server.shutdown()
        server.server_close()


def test_execute_missing_env_and_server_count_mismatch(synthetic_json_path: Path, monkeypatch):
    monkeypatch.delenv("SUPABASE_URL", raising=False)
    monkeypatch.delenv("SUPABASE_SERVICE_ROLE_KEY", raising=False)
    with pytest.raises(ValueError, match="SUPABASE_URL"):
        mig.main([str(synthetic_json_path), "--execute"])

    # Giả lập server trả thiếu 1 dòng khi đối soát sau nạp -> phải raise ValueError
    class MismatchHandler(BaseHTTPRequestHandler):
        def log_message(self, fmt, *args):
            pass

        def do_GET(self):
            parsed = urlparse(self.path)
            body = json.dumps({"users": []}).encode("utf-8") if parsed.path == "/auth/v1/admin/users" else b"[]"
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            if parsed.path.startswith("/rest/v1/"):
                self.send_header("Content-Range", "*/0")  # Cố tình báo 0 dòng trên server
            self.end_headers()
            self.wfile.write(body)

        def do_POST(self):
            length = int(self.headers.get("Content-Length", "0"))
            if length:
                self.rfile.read(length)
            body = json.dumps({"id": "00000000-0000-0000-0000-000000000001"}).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def do_DELETE(self):
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", "2")
            self.end_headers()
            self.wfile.write(b"[]")

    server = ThreadingHTTPServer(("127.0.0.1", 0), MismatchHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        monkeypatch.setenv("SUPABASE_URL", f"http://127.0.0.1:{server.server_port}")
        monkeypatch.setenv("SUPABASE_SERVICE_ROLE_KEY", "test-key")
        monkeypatch.delenv("INITIAL_USER_PASSWORD", raising=False)
        with pytest.raises(ValueError, match="Đối soát trên server thất bại"):
            mig.main([str(synthetic_json_path), "--execute"])
    finally:
        server.shutdown()
        server.server_close()


# ---------------------------------------------------------------------------
# 6. Bộ 5 kiểm thử hồi quy vòng QA & Critic (UTC, Clean Seed/FK, UUID, Flat JSON, Env/Ban)
# ---------------------------------------------------------------------------


def test_blocker1_utc_session_no_plus7h_shift_sql_and_execute(tmp_path: Path, monkeypatch):
    """Kiểm tra chống nhảy +7 giờ khi Postgres chạy ở UTC (cả --sql-out lẫn --execute)."""
    db_dict = _make_synthetic_db_dict()
    # Đặt mốc 18:30 tối giờ VN: nếu bị hiểu nhầm là 18:30 UTC thì AT TIME ZONE 'Asia/Ho_Chi_Minh' nhảy sang 01:30 ngày hôm sau!
    db_dict["tabs"]["lines"]["rows"][0][21] = "2026-03-15T08:15:00"  # created_at
    db_dict["tabs"]["lines"]["rows"][0][22] = "2026-03-15T18:30:00"  # updated_at
    db_dict["tabs"]["requests"]["rows"][1][4] = "2026-03-15T18:30:00"  # issued_at
    db_dict["tabs"]["receipts"]["rows"][0][6] = "2026-03-15T18:30:00"  # created_at
    db_dict["tabs"]["receipts"]["rows"][1][6] = "2026-03-15T18:30:00"  # created_at
    db_dict["tabs"]["audit_log"]["rows"] = [
        ["1", "2026-03-15T18:30:00", "admin", "lines", "10", "note", "", "sửa"]
    ]

    src_path = tmp_path / "utc_check.json"
    src_path.write_text(json.dumps(db_dict, ensure_ascii=False), encoding="utf-8")
    sql_path = tmp_path / "utc_cutover.sql"

    mig.main([str(src_path), "--sql-out", str(sql_path)])
    sql_text = sql_path.read_text(encoding="utf-8")
    assert "BEGIN;\nSET LOCAL TIME ZONE 'Asia/Ho_Chi_Minh';" in sql_text

    pglite_dir = WORKTREE_ROOT / "node_modules" / "@electric-sql" / "pglite"
    shim_sql = WORKTREE_ROOT / "tests" / "helpers" / "supabase_local_shim.sql"
    schema_sql = WORKTREE_ROOT / "supabase" / "migrations" / "001_schema_rls_rpc.sql"
    assert pglite_dir.exists(), "Cần @electric-sql/pglite để kiểm chứng UTC"

    js_runner = tmp_path / "verify_utc.js"
    js_runner.write_text(
        f"""
        const fs = require('fs');
        const {{ PGlite }} = require({json.dumps(str(pglite_dir))});
        (async () => {{
          const db = new PGlite();
          if ({json.dumps(shim_sql.exists() and schema_sql.exists())}) {{
            await db.exec(fs.readFileSync({json.dumps(str(shim_sql))}, 'utf8'));
            await db.exec(fs.readFileSync({json.dumps(str(schema_sql))}, 'utf8'));
          }} else {{
            await db.exec(`
              CREATE SCHEMA IF NOT EXISTS private;
              CREATE OR REPLACE FUNCTION private.fmt_ts(p_ts timestamptz) RETURNS text LANGUAGE sql STABLE AS $$
                SELECT CASE WHEN p_ts IS NULL THEN NULL ELSE to_char(p_ts AT TIME ZONE 'Asia/Ho_Chi_Minh', 'YYYY-MM-DD"T"HH24:MI:SS') END;
              $$;
              CREATE TABLE public.suppliers (id bigint generated by default as identity primary key, code text, name text, address text, phone text, email text, tax_code text, field text);
              CREATE TABLE public.project_items (id bigint generated by default as identity primary key, project_code text, project_name text, item_code text, item_name text, owner text, status text);
              CREATE TABLE public.material_groups (id bigint generated by default as identity primary key, name text, description text);
              CREATE TABLE public.units (id bigint generated by default as identity primary key, name text);
              CREATE TABLE public.uoms (id bigint generated by default as identity primary key, name text);
              CREATE TABLE public.signers (id bigint generated by default as identity primary key, unit_contains text, position text, title text, name text);
              CREATE TABLE public.requests (id bigint generated by default as identity primary key, code text unique not null, year int not null, seq int not null, issued_at timestamptz, issued_by text, header_json jsonb not null, line_count int not null default 0, legacy_flag boolean not null default false);
              CREATE TABLE public.lines (id bigint generated by default as identity primary key, request_code text, unit text, material_group text, name text, uom text, qty_requested numeric, qty_on_hand numeric, need_date date, project_code text, project_name text, purpose text, equipment_name text, delivery_location text, requester text, request_date date, supplier text, unit_price numeric, note text, legacy_flag boolean not null default false, created_by text, created_at timestamptz, updated_at timestamptz, qty_received numeric not null default 0, last_received_date date);
              CREATE TABLE public.receipts (id bigint generated by default as identity primary key, line_id bigint not null references public.lines(id), received_date date, qty numeric not null, received_by text, note text, created_at timestamptz, created_by uuid);
              CREATE TABLE public.audit_log (id bigint generated by default as identity primary key, at timestamptz, "user" text, "table" text, row_id text, field text, "old" text, "new" text);
              CREATE TABLE public.sequences (id bigint generated by default as identity primary key, key text unique not null, last_value bigint not null);
            `);
          }}
          // Ép phiên chạy ở múi giờ UTC giống Supabase Cloud trước khi chạy file SQL Cutover!
          await db.exec("SET TIME ZONE 'UTC';");
          await db.exec(fs.readFileSync({json.dumps(str(sql_path))}, 'utf8'));

          const lineTs = (await db.query("SELECT private.fmt_ts(created_at) AS c, private.fmt_ts(updated_at) AS u FROM public.lines WHERE id = 10")).rows[0];
          const reqTs = (await db.query("SELECT private.fmt_ts(issued_at) AS i FROM public.requests WHERE id = 2")).rows[0];
          const recTs = (await db.query("SELECT private.fmt_ts(created_at) AS r FROM public.receipts WHERE id = 1")).rows[0];
          const audTs = (await db.query("SELECT private.fmt_ts(at) AS a FROM public.audit_log WHERE id = 1")).rows[0];
          await db.close();
          console.log(JSON.stringify({{ lineTs, reqTs, recTs, audTs }}));
        }})().catch(e => {{ console.error(e); process.exit(1); }});
        """,
        encoding="utf-8",
    )
    proc = subprocess.run(["node", str(js_runner)], capture_output=True, text=True, check=True, encoding="utf-8")
    out = json.loads(proc.stdout.strip().splitlines()[-1])
    assert out["lineTs"] == {"c": "2026-03-15T08:15:00", "u": "2026-03-15T18:30:00"}
    assert out["reqTs"]["i"] == "2026-03-15T18:30:00"
    assert out["recTs"]["r"] == "2026-03-15T18:30:00"
    assert out["audTs"]["a"] == "2026-03-15T18:30:00"

    # Kiểm tra --execute gắn hậu tố +07:00 rõ ràng trong JSON payload gửi lên PostgREST
    posted_rows: dict[str, list[dict]] = {}

    class CaptureHandler(BaseHTTPRequestHandler):
        def log_message(self, fmt, *args):
            pass

        def _send(self, code: int, data, hdrs: dict | None = None):
            b = json.dumps(data).encode("utf-8")
            self.send_response(code)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(b)))
            for k, v in (hdrs or {}).items():
                self.send_header(k, v)
            self.end_headers()
            self.wfile.write(b)

        def do_GET(self):
            parsed = urlparse(self.path)
            if parsed.path == "/auth/v1/admin/users":
                self._send(200, {"users": []})
                return
            tbl = parsed.path.split("/rest/v1/")[1]
            cnt = len(posted_rows.get(tbl, []))
            self._send(200, [], {"Content-Range": f"0-0/{cnt}" if cnt else "*/0"})

        def do_DELETE(self):
            parsed = urlparse(self.path)
            tbl = parsed.path.split("/rest/v1/")[1]
            posted_rows[tbl] = []
            self._send(200, [])

        def do_POST(self):
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length).decode("utf-8")) if length else None
            parsed = urlparse(self.path)
            if parsed.path == "/auth/v1/admin/users":
                uid = f"00000000-0000-0000-0000-{len(posted_rows.get('_auth', [])) + 1:012d}"
                posted_rows.setdefault("_auth", []).append(payload)
                self._send(200, {"id": uid, "email": payload["email"]})
                return
            tbl = parsed.path.split("/rest/v1/")[1]
            posted_rows.setdefault(tbl, []).extend(payload if isinstance(payload, list) else [payload])
            self._send(201, [])

    server = ThreadingHTTPServer(("127.0.0.1", 0), CaptureHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        monkeypatch.setenv("SUPABASE_URL", f"http://127.0.0.1:{server.server_port}")
        monkeypatch.setenv("SUPABASE_SERVICE_ROLE_KEY", "srv-key")
        monkeypatch.setenv("INITIAL_USER_PASSWORD", "MatKhauTam@2026")
        mig.main([str(src_path), "--execute"])

        l10_rest = next(r for r in posted_rows["lines"] if r["id"] == 10)
        assert l10_rest["created_at"] == "2026-03-15T08:15:00+07:00"
        assert l10_rest["updated_at"] == "2026-03-15T18:30:00+07:00"
        req2_rest = next(r for r in posted_rows["requests"] if r["id"] == 2)
        assert req2_rest["issued_at"] == "2026-03-15T18:30:00+07:00"
        req1_rest = next(r for r in posted_rows["requests"] if r["id"] == 1)
        assert req1_rest["issued_at"] is None
        assert posted_rows["receipts"][0]["created_at"] == "2026-03-15T18:30:00+07:00"
        assert posted_rows["audit_log"][0]["at"] == "2026-03-15T18:30:00+07:00"
    finally:
        server.shutdown()
        server.server_close()


def test_blocker2_execute_cleans_seed_data_and_sanitizes_receipt_created_by(tmp_path: Path, monkeypatch):
    """Kiểm tra --execute xóa dữ liệu seed/test cũ theo thứ tự ngược FK và lọc sạch receipts.created_by."""
    db_dict = _make_synthetic_db_dict()
    valid_admin_uuid = "00000000-0000-0000-0000-000000000001"
    unknown_uuid = "deadbeef-dead-beef-dead-beefdeadbeef"
    db_dict["tabs"]["receipts"]["headers"].append("created_by")
    db_dict["tabs"]["receipts"]["rows"][0].append("nam")             # tên user cũ (không phải UUID) -> None
    db_dict["tabs"]["receipts"]["rows"][1].append(unknown_uuid)      # UUID lạ không có trong auth.users -> None
    db_dict["tabs"]["receipts"]["rows"][2].append(valid_admin_uuid)  # UUID hợp lệ khớp admin -> giữ nguyên

    src_path = tmp_path / "dirty_seed.json"
    src_path.write_text(json.dumps(db_dict, ensure_ascii=False), encoding="utf-8")

    # Kiểm tra transform_data lọc ngay chuỗi không phải UUID ("nam" -> None)
    raw = mig.load_source(src_path)
    bundle = mig.transform_data(raw)
    assert bundle.tables["receipts"][0]["created_by"] is None
    # Và generate_sql không để lọt UUID chưa xác thực vào FK receipts_created_by_fkey
    sql_out = mig.generate_sql(bundle)
    assert unknown_uuid not in sql_out

    # Giả lập Supabase đang chứa sẵn dữ liệu rác từ lần chạy thử trước
    delete_calls: list[tuple[str, str]] = []
    server_rows: dict[str, list[dict]] = {
        "signers": [{"id": 999, "title": "Seed cũ dư thừa"}],
        "lines": [{"id": 88888, "name": "Dòng test cũ"}],
        "profiles": [{"id": "99999999-9999-9999-9999-999999999999", "username": "ghost"}],
    }
    auth_users = [{"id": valid_admin_uuid, "email": "admin@dncvt.local"}]

    class SeedCleanupHandler(BaseHTTPRequestHandler):
        def log_message(self, fmt, *args):
            pass

        def _send(self, code: int, data, hdrs: dict | None = None):
            b = json.dumps(data).encode("utf-8")
            self.send_response(code)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(b)))
            for k, v in (hdrs or {}).items():
                self.send_header(k, v)
            self.end_headers()
            self.wfile.write(b)

        def do_GET(self):
            parsed = urlparse(self.path)
            if parsed.path == "/auth/v1/admin/users":
                self._send(200, {"users": auth_users})
                return
            tbl = parsed.path.split("/rest/v1/")[1]
            cnt = len(server_rows.get(tbl, []))
            self._send(200, [], {"Content-Range": f"0-0/{cnt}" if cnt else "*/0"})

        def do_DELETE(self):
            parsed = urlparse(self.path)
            tbl = parsed.path.split("/rest/v1/")[1]
            delete_calls.append((tbl, parsed.query))
            qs = parse_qs(parsed.query)
            if tbl == "profiles" and "id" in qs and qs["id"][0].startswith("not.in."):
                keep_raw = qs["id"][0][len("not.in.(") : -1]
                keep_set = {x.strip().strip('"') for x in keep_raw.split(",") if x.strip()}
                server_rows["profiles"] = [r for r in server_rows.get("profiles", []) if r["id"] in keep_set]
            else:
                server_rows[tbl] = []
            self._send(200, [])

        def do_POST(self):
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length).decode("utf-8")) if length else None
            parsed = urlparse(self.path)
            if parsed.path == "/auth/v1/admin/users":
                uid = f"00000000-0000-0000-0000-{len(auth_users) + 1:012d}"
                u = {"id": uid, "email": payload["email"]}
                auth_users.append(u)
                self._send(200, u)
                return
            tbl = parsed.path.split("/rest/v1/")[1]
            items = payload if isinstance(payload, list) else [payload]
            existing_by_id = {r["id"]: r for r in server_rows.get(tbl, [])}
            for it in items:
                existing_by_id[it["id"]] = it
            server_rows[tbl] = list(existing_by_id.values())
            self._send(201, [])

    server = ThreadingHTTPServer(("127.0.0.1", 0), SeedCleanupHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        monkeypatch.setenv("SUPABASE_URL", f"http://127.0.0.1:{server.server_port}")
        monkeypatch.setenv("SUPABASE_SERVICE_ROLE_KEY", "srv-key")
        monkeypatch.setenv("INITIAL_USER_PASSWORD", "MatKhauTam@2026")

        mig.main([str(src_path), "--execute"])

        # Kiểm tra thứ tự DELETE các bảng dữ liệu đúng thứ tự ngược khóa ngoại
        deleted_data_tables = [tbl for tbl, _ in delete_calls if tbl in mig.SQL_ID_TABLES]
        assert deleted_data_tables == list(reversed(mig.SQL_ID_TABLES))
        assert any(tbl == "profiles" for tbl, _ in delete_calls)

        # Kiểm tra dữ liệu rác đã sạch và receipts.created_by chỉ giữ UUID hợp lệ có thật trong auth
        assert len(server_rows["profiles"]) == 2
        assert all(p["username"] != "ghost" for p in server_rows["profiles"])
        rec_by_id = {r["id"]: r for r in server_rows["receipts"]}
        assert rec_by_id[1]["created_by"] is None
        assert rec_by_id[2]["created_by"] is None
        assert rec_by_id[3]["created_by"] == valid_admin_uuid
    finally:
        server.shutdown()
        server.server_close()


def test_seam01_supabase_export_xlsx_with_uuid_users_and_audit(tmp_path: Path):
    """Kiểm tra vòng khép kín: xuất .xlsx từ Supabase (users.id kiểu UUID) -> migrate_to_supabase.py --dry-run."""
    db_dict = _make_synthetic_db_dict()
    xlsx_path = tmp_path / "ncvt-toan-bo-supabase.xlsx"
    _write_admin_export_xlsx(db_dict, xlsx_path, include_seeded_signers=True)

    # Ghi đè sheet users và audit_log giống hệt dữ liệu xuất từ rpc_export_all của Supabase
    wb = openpyxl.load_workbook(xlsx_path)
    ws_users = wb["users"]
    ws_users.delete_rows(1, ws_users.max_row)
    ws_users.append(["id", "username", "full_name", "role", "active"])
    ws_users.append(["ab8dcaa9-10a3-4860-8e7f-292e23657e5f", "admin", "Quản trị", "admin", True])
    ws_users.append(["cd9efbb0-21b4-4971-9f8a-303f34768f6a", "nam", "Nguyễn Văn Nam", "user", False])

    ws_audit = wb["audit_log"]
    ws_audit.delete_rows(1, ws_audit.max_row)
    ws_audit.append(["id", "at", "user", "table", "record_id", "field", "old", "new"])
    ws_audit.append([1, "2026-10-26T10:00:00", "admin", "users", "ab8dcaa9-10a3-4860-8e7f-292e23657e5f", "role", "user", "admin"])
    wb.save(xlsx_path)

    bundle = mig.main([str(xlsx_path), "--dry-run"])
    sql_text = mig.generate_sql(bundle)

    assert bundle.tables["users"][0]["id"] == "ab8dcaa9-10a3-4860-8e7f-292e23657e5f"
    assert bundle.tables["users"][1]["active"] is False
    assert bundle.tables["audit_log"][0]["row_id"] == "ab8dcaa9-10a3-4860-8e7f-292e23657e5f"
    assert "ab8dcaa9-10a3-4860-8e7f-292e23657e5f" in sql_text


def test_seam02_flat_json_from_rpc_export_all_and_backup_snapshot(tmp_path: Path):
    """Kiểm tra _load_json_source đọc được cả JSON phẳng từ rpc_export_all / private.backup_snapshots.payload."""
    flat_payload = {
        "lines": [
            {
                "id": 10,
                "request_code": "YC2026-294",
                "unit": "P.Cải tiến",
                "material_group": "C",
                "name": "Thép tấm 10ly",
                "uom": "tấm",
                "qty_requested": 4,
                "qty_on_hand": 0,
                "need_date": "2026-10-30",
                "project_code": "Ms027",
                "project_name": "Hệ thống Lò NL",
                "purpose": "Chế tạo",
                "equipment_name": None,
                "delivery_location": "Kho 1",
                "requester": "Hà",
                "request_date": "2026-10-25",
                "supplier": "KHVT",
                "unit_price": 250000,
                "note": None,
                "legacy_flag": False,
                "created_by": "ha",
                "created_at": "2026-10-25T08:00:00",
                "updated_at": "2026-10-25T09:00:00",
            }
        ],
        "requests": [
            {
                "id": 1,
                "code": "YC2026-294",
                "year": 2026,
                "seq": 294,
                "issued_at": "2026-10-25T08:30:00",
                "issued_by": "Hà",
                "header_json": {"unit": "P.Cải tiến", "project_code": "Ms027", "note": ""},
                "line_count": 1,
                "legacy_flag": False,
            }
        ],
        "receipts": [
            {
                "id": 1,
                "line_id": 10,
                "received_date": "2026-10-26",
                "qty": 2,
                "received_by": "Hà",
                "note": None,
                "created_at": "2026-10-26T10:00:00",
            }
        ],
        "suppliers": [{"id": 1, "code": "KHVT", "name": "Phòng KHVT"}],
        "project_items": [{"id": 27, "project_code": "Ms027", "item_name": "Hệ thống Lò NL"}],
        "material_groups": [{"id": 1, "name": "C", "description": "Thiết bị"}],
        "signers": [],
        "units": [{"id": 1, "name": "P.Cải tiến"}],
        "uoms": [{"id": 1, "name": "tấm"}],
        "users": [
            {"id": "ab8dcaa9-10a3-4860-8e7f-292e23657e5f", "username": "admin", "full_name": "Quản trị", "role": "admin", "active": True}
        ],
        "audit_log": [],
        "sequences": [{"id": 1, "key": "YC2026", "last_value": 294}],
    }

    flat_path = tmp_path / "export_all.json"
    flat_path.write_text(json.dumps(flat_payload, ensure_ascii=False), encoding="utf-8")
    b_flat = mig.main([str(flat_path), "--dry-run"])
    assert len(b_flat.tables["lines"]) == 1
    assert b_flat.tables["lines"][0]["project_code"] == "MS027"
    assert b_flat.tables["lines"][0]["qty_received"] == 2
    assert b_flat.tables["users"][0]["id"] == "ab8dcaa9-10a3-4860-8e7f-292e23657e5f"

    # Cũng đọc được khi bọc trong {"payload": ...} của private.backup_snapshots
    snap_path = tmp_path / "snapshot_row.json"
    snap_path.write_text(json.dumps({"name": "ncvt-backup", "payload": flat_payload}, ensure_ascii=False), encoding="utf-8")
    b_snap = mig.main([str(snap_path), "--dry-run"])
    assert b_snap.tables["lines"] == b_flat.tables["lines"]


def test_minor_synthetic_email_domain_and_inactive_ban_duration(tmp_path: Path, monkeypatch):
    """Kiểm tra SYNTHETIC_EMAIL_DOMAIN tùy chỉnh và ban_duration='876000h' cho user active=False."""
    db_dict = _make_synthetic_db_dict()
    db_dict["tabs"]["users"]["rows"] = [
        ["1", "admin", "Quản trị", "admin", "h1", "s1", "TRUE", "2026-10-23T09:13:54"],
        ["2", "khoa.nv", "Nguyễn Văn Khóa", "user", "h2", "s2", "", "2026-10-23T09:14:03"],  # active="" -> False trong db.json
    ]
    src_path = tmp_path / "users_ban.json"
    src_path.write_text(json.dumps(db_dict, ensure_ascii=False), encoding="utf-8")

    created_auth_payloads: list[dict] = []
    server_rows: dict[str, list[dict]] = {}

    class AuthInspectHandler(BaseHTTPRequestHandler):
        def log_message(self, fmt, *args):
            pass

        def _send(self, code: int, data, hdrs: dict | None = None):
            b = json.dumps(data).encode("utf-8")
            self.send_response(code)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(b)))
            for k, v in (hdrs or {}).items():
                self.send_header(k, v)
            self.end_headers()
            self.wfile.write(b)

        def do_GET(self):
            parsed = urlparse(self.path)
            if parsed.path == "/auth/v1/admin/users":
                self._send(200, {"users": []})
                return
            tbl = parsed.path.split("/rest/v1/")[1]
            cnt = len(server_rows.get(tbl, []))
            self._send(200, [], {"Content-Range": f"0-0/{cnt}" if cnt else "*/0"})

        def do_DELETE(self):
            parsed = urlparse(self.path)
            tbl = parsed.path.split("/rest/v1/")[1]
            server_rows[tbl] = []
            self._send(200, [])

        def do_POST(self):
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length).decode("utf-8")) if length else None
            parsed = urlparse(self.path)
            if parsed.path == "/auth/v1/admin/users":
                created_auth_payloads.append(payload)
                uid = f"00000000-0000-0000-0000-{len(created_auth_payloads):012d}"
                self._send(200, {"id": uid, "email": payload["email"]})
                return
            tbl = parsed.path.split("/rest/v1/")[1]
            server_rows.setdefault(tbl, []).extend(payload if isinstance(payload, list) else [payload])
            self._send(201, [])

    server = ThreadingHTTPServer(("127.0.0.1", 0), AuthInspectHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        monkeypatch.setenv("SUPABASE_URL", f"http://127.0.0.1:{server.server_port}")
        monkeypatch.setenv("SUPABASE_SERVICE_ROLE_KEY", "srv-key")
        monkeypatch.setenv("INITIAL_USER_PASSWORD", "MatKhauTam@2026")
        monkeypatch.setenv("SYNTHETIC_EMAIL_DOMAIN", "pct.internal")

        mig.main([str(src_path), "--execute"])

        assert [p["email"] for p in created_auth_payloads] == ["admin@pct.internal", "khoa.nv@pct.internal"]
        assert "ban_duration" not in created_auth_payloads[0]
        assert created_auth_payloads[1].get("ban_duration") == "876000h"
        prof_by_user = {p["username"]: p for p in server_rows["profiles"]}
        assert prof_by_user["admin"]["active"] is True
        assert prof_by_user["khoa.nv"]["active"] is False
    finally:
        server.shutdown()
        server.server_close()
