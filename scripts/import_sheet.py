#!/usr/bin/env python3
"""
Import dữ liệu từ Google Sheet cũ (file .xlsx tải về) sang cấu trúc của app.

    python scripts/import_sheet.py samples/de_nghi_cu.xlsx            # tạo out/db_khoi_tao.xlsx
    python scripts/import_sheet.py samples/de_nghi_cu.xlsx --dry-run  # chỉ in báo cáo đối soát

Kết quả (thư mục out/, ghi đè mỗi lần chạy → chạy lại không nhân đôi dữ liệu):
  - db_khoi_tao.xlsx   : các tab lines, requests, receipts, danh mục, sequences — tải lên Google Drive,
                         mở bằng Google Sheets, gắn Apps Script rồi chạy setup() (xem README).
  - import_errors.xlsx : dòng, cột, giá trị, lý do — các giá trị không hiểu / không khớp danh mục (không tự đoán).
  - import_report.txt  : báo cáo đối soát (tổng dòng, tổng thành tiền, số phiếu theo tháng, số dòng lỗi).

File nguồn chỉ được ĐỌC, không bao giờ bị sửa.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import unicodedata
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
SCHEMA_JS = ROOT / "apps-script" / "Schema.js"
INITIAL_SEQ = {2026: 294}
CODE_RE = re.compile(r"^YC(\d{4})-(\d+)$")

# ---------------------------------------------------------------- tiện ích chung


def load_schema(path: Path = SCHEMA_JS) -> dict[str, list[str]]:
    """Đọc danh sách cột từ apps-script/Schema.js để luôn khớp với backend."""
    text = path.read_text(encoding="utf-8")
    body = text[text.index("const SCHEMA"):text.index("};", text.index("const SCHEMA"))]
    schema = {}
    for m in re.finditer(r"(\w+):\s*\[((?:\s*\['\w+',\s*'\w+'\],?\s*)+)\]", body):
        schema[m.group(1)] = re.findall(r"\['(\w+)',\s*'\w+'\]", m.group(2))
    if "lines" not in schema or "requests" not in schema:
        raise SystemExit(f"Không đọc được SCHEMA trong {path}")
    return schema


def fold(s) -> str:
    """Bỏ dấu, chữ thường, gộp ký tự lạ thành 1 khoảng trắng: 'Nhà cc/Người cung cấp' -> 'nha cc nguoi cung cap'."""
    s = unicodedata.normalize("NFD", str(s or "")).replace("đ", "d").replace("Đ", "D")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn").lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def is_empty(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def clean_text(v) -> str | None:
    if is_empty(v):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (dt.datetime, dt.date)):
        v = v.strftime("%d/%m/%Y")
    # NFC: Sheet cũ lẫn 2 kiểu mã hóa dấu ("P.Cải tiến" dựng sẵn / tổ hợp) → gộp về một
    s = re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(v))).strip()
    return s or None


def parse_num(v):
    """Trả (số|None, lỗi|None). Hiểu '1.234,5', '1,234.5', '1,5', '1.500' (=1500), '1.500.000'.
    Lỗi bắt đầu bằng 'Cảnh báo:' vẫn trả kèm giá trị."""
    if is_empty(v):
        return None, None
    if isinstance(v, bool):
        return None, "Không phải số"
    if isinstance(v, (int, float)):
        return float(v), None
    s = re.sub(r"\s", "", str(v))
    last_dot, last_comma = s.rfind("."), s.rfind(",")
    if last_dot >= 0 and last_comma >= 0:
        s = s.replace(".", "").replace(",", ".") if last_comma > last_dot else s.replace(",", "")
    elif last_comma >= 0:
        s = s.replace(",", "") if re.fullmatch(r"\d{1,3}(,\d{3})+", s) else s.replace(",", ".")
    elif re.fullmatch(r"[1-9]\d{0,2}(\.\d{3})+", s):
        # Định dạng Việt Nam: dấu chấm phân cách nghìn
        raw, s = s, s.replace(".", "")
        if raw.count(".") == 1:
            return float(s), f"Cảnh báo: '{raw}' được hiểu là {s} (dấu chấm phân cách nghìn) — kiểm tra lại"
    try:
        return float(s), None
    except ValueError:
        return None, "Không phải số"


MONTHS_EN = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def parse_date(v, ref: dt.date | None = None):
    """
    Trả (date|None, lỗi|None). Nhận: datetime, số serial Excel, '23/12/2021', '23-12-2021', '23.12.21',
    '2021-12-23', và dạng thiếu năm '24-thg 12' / '11-Thg3' / '24/12': chọn năm để ngày gần `ref` nhất
    (ref = Ngày yêu cầu khi đọc ngày cần/ngày nhận; = Ngày cần có khi chính Ngày yêu cầu thiếu năm).
    """
    if is_empty(v):
        return None, None
    if isinstance(v, dt.datetime):
        return v.date(), None
    if isinstance(v, dt.date):
        return v, None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 20000 < v < 80000:
            return dt.date(1899, 12, 30) + dt.timedelta(days=int(v)), None
        return None, "Số không phải ngày"
    s = str(v).strip().lower()
    if re.fullmatch(r"\d{5}(\.0+)?", s):
        return parse_date(float(s), ref)
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[ t].*)?", s)
    if m:
        return _mk(int(m[1]), int(m[2]), int(m[3]))
    m = re.fullmatch(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2}|\d{4})", s)
    if m:
        y = int(m[3]) + (2000 if len(m[3]) == 2 else 0)
        return _mk(y, int(m[2]), int(m[1]))
    # Thiếu năm: '24-thg 12', '24 thg 12', '24/12', '24-dec'
    m = re.fullmatch(r"(\d{1,2})\s*[/\-.\s]\s*(?:thg|th|thang)?\s*(\d{1,2})", s) or \
        re.fullmatch(r"(\d{1,2})\s*[\-\s]\s*([a-z]{3})[a-z]*", s)
    if m:
        month = int(m[2]) if m[2].isdigit() else MONTHS_EN.get(m[2])
        if not month:
            return None, "Không hiểu ngày"
        if ref is None:
            return None, "Ngày thiếu năm và không có ngày tham chiếu để suy ra"
        # vd. yêu cầu 20/12/2021, cần '05-thg 1' → 05/01/2022 (gần hơn 05/01/2021)
        options = [_mk(y, month, int(m[1])) for y in (ref.year - 1, ref.year, ref.year + 1)]
        valid = [o for o in options if o[0]]
        if not valid:
            return options[1]
        return min(valid, key=lambda o: abs((o[0] - ref).days))
    return None, "Không hiểu ngày"


def _mk(y, mo, d):
    try:
        return dt.date(y, mo, d), None
    except ValueError:
        return None, "Ngày không tồn tại"


def iso(d: dt.date | None) -> str | None:
    return d.isoformat() if d else None


def num_out(x):
    if x is None:
        return None
    return int(x) if float(x).is_integer() else round(x, 6)


# ---------------------------------------------------------------- nhận diện cột

DATA_HEADERS = {
    "stt": ["stt", "stt tu dong"],
    "unit": ["don vi de xuat"],
    "material_group": ["nhom vt", "nhom vat tu"],
    "name": ["ten vat tu"],
    "uom": ["dvt", "don vi tinh"],
    "qty_requested": ["so luong yeu cau", "sl yeu cau"],
    "qty_on_hand": ["sl ton kho", "so luong ton kho", "ton kho"],
    "need_date": ["ngay can co", "ngay can"],
    "project_code": ["ma hang muc", "ma hm"],
    "project_name": ["hang muc"],
    "purpose": ["muc dich su dung", "muc dich"],
    "delivery_location": ["dia diem giao", "dia diem giao hang", "noi giao"],
    "supplier": ["nha cc nguoi mua", "nha cc nguoi cung cap", "nha cung cap nguoi cung cap", "nha cung cap", "nguoi cung cap", "nha cc"],
    "request_date": ["ngay yeu cau"],
    "received_date": ["ngay nhan"],
    "qty_received": ["so luong thuc nhan", "sl thuc nhan"],
    "requester": ["nguoi yeu cau"],
    "signed_by": ["ky nhan"],
    "status": ["trang thai"],
    "request_code": ["so phieu"],
    "days_left": ["so ngay con lai"],
    "unit_price": ["don gia"],
    "amount": ["thanh tien"],
    "equipment_name": ["ten thiet bi", "ten thiet bi hang muc"],
    "work_code": ["pl ma cv", "phan loai ma cv"],
    "week": ["tuan"],
    "month": ["thang"],
    "qty_missing": ["so luong con thieu", "sl con thieu"],
}
MULTI_COLUMNS = {"purpose"}  # "Mục đích sử dụng" + "Mục đích sử dụng2" (cột 2 = "Hạng mục-Mục đích", chỉ dùng khi cột 1 trống)

SUPPLIER_HEADERS = {
    "code": ["mancc", "ma ncc", "ma nha cung cap", "ma", "ten viet tat", "viet tat"],
    "name": ["nhacc", "ten ncc", "ten nha cung cap", "nha cung cap", "ten", "ten cong ty", "cong ty"],
    "address": ["diachincc", "dia chi"],
    "phone": ["sdtlh", "sdt", "so dien thoai", "dien thoai", "dt"],
    "email": ["dcemail", "email", "mail"],
    "tax_code": ["mskd", "mst", "ma so thue"],
    "field": ["linhvucchuyenmon", "linh vuc", "nganh hang", "nganh", "mat hang"],
}
PROJECT_HEADERS = {
    "project_code": ["ma da", "ma du an", "ma du an mdm"],
    "project_name": ["ten da", "ten du an", "du an", "ten du an mdm"],
    "item_code": ["ma hm", "ma hang muc", "ma hang muc mdm"],
    "item_name": ["ten hm", "ten hang muc", "hang muc", "hang muc mdm"],
    "owner": ["phu trach", "nguoi phu trach", "chu tri", "pm"],
    "status": ["trang thai", "tinh trang"],
}
LEGACY_GROUP_ALIASES = {
    "vt che tao": "C",
    "bulong": "C",
    "gcck": "C",
    "vbi": "C",
    "cat ton": "C",
    "cong cu ca nhan": "C",
    "thiet bi": "C",
    "van phong pham": "E",
    "vt van phong": "E",
    "do gia dung": "E",
}


def find_sheet(wb, *names):
    wanted = [fold(n) for n in names]
    for ws in wb.worksheets:
        if fold(ws.title) in wanted:
            return ws
    return None


def detect_header(rows, aliases, must_have, scan=15):
    """Tìm dòng tiêu đề trong `scan` dòng đầu; trả (chỉ số dòng, {field: [cột,...]}, [cột lạ])."""
    lookup = {a: f for f, al in aliases.items() for a in al}
    for i, row in enumerate(rows[:scan]):
        names = [fold(c) for c in row]
        found = [lookup.get(n) for n in names]
        if not any(m in found for m in (must_have if isinstance(must_have, tuple) else (must_have,))):
            continue
        mapping: dict[str, list[int]] = defaultdict(list)
        unknown = []
        for j, n in enumerate(names):
            if not n:
                continue
            f = lookup.get(n) or lookup.get(re.sub(r"\s*\d+$", "", n))  # 'muc dich su dung 2'
            if f and (f in MULTI_COLUMNS or not mapping[f]):
                mapping[f].append(j)
            else:
                unknown.append(str(row[j]))
        return i, dict(mapping), unknown
    return None, {}, []


def read_rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


# ---------------------------------------------------------------- import


class Importer:
    def __init__(self, src: Path, pending_days: int = 30):
        self.src = src
        self.pending_days = pending_days
        self.schema = load_schema()
        self.errors: list[dict] = []
        self.notes: list[str] = []

    def err(self, row, column, value, reason):
        self.errors.append({"Dòng": row, "Cột": column, "Giá trị": "" if value is None else str(value), "Lý do": reason})

    # ---- danh mục
    def read_catalog(self, ws, aliases, must_have):
        if ws is None:
            return []
        rows = read_rows(ws)
        hi, _, _ = detect_header(rows, aliases, must_have)
        if hi is None:
            self.notes.append(f'Tab "{ws.title}": không tìm thấy dòng tiêu đề, bỏ qua.')
            return []
        # Một tab có thể chứa nhiều bảng đặt cạnh nhau (vd. QL Ma DAHM): mỗi bảng bắt đầu bằng cột "STT"
        header = rows[hi]
        starts = [j for j, c in enumerate(header) if fold(c) == "stt"]
        starts = sorted(set([0] + starts))
        bounds = list(zip(starts, starts[1:] + [len(header)]))
        out, seen = [], set()
        for a, b in bounds:
            _, mapping, _ = detect_header([header[a:b]], aliases, must_have, scan=1)
            if not mapping:
                continue
            self.notes.append(f'Tab "{ws.title}" (cột {a + 1}–{b}): nhận ' + ", ".join(f"{f}←{header[a + c[0]]}" for f, c in mapping.items()))
            for r in rows[hi + 1:]:
                seg = list(r[a:b])
                rec = {f: clean_text(seg[c[0]]) if c[0] < len(seg) else None for f, c in mapping.items()}
                key = tuple(rec.values())
                if any(rec.values()) and key not in seen:
                    seen.add(key)
                    out.append(rec)
        return out

    def read_material_groups(self, ws):
        """Tab Nhóm VT: cột A là mã nhóm (C, E… có thể gộp ô nhiều dòng), cột B+ là giải thích."""
        if ws is None:
            return []
        grouped: OrderedDict[str, list[str]] = OrderedDict()
        current = None
        for r in read_rows(ws):
            code = clean_text(r[0]) if r else None
            if code and fold(code) not in ("nhom vt", "nhom vat tu", "stt") and not re.fullmatch(r"\d+", code):
                current = code
                grouped.setdefault(current, [])
            if current:
                for v in r[1:]:
                    desc = clean_text(v)
                    if desc:
                        part = desc.rstrip(".;").strip()
                        if part and part not in grouped[current]:
                            grouped[current].append(part)
        return [{"name": k, "description": "; ".join(v) if v else None} for k, v in grouped.items()]

    # ---- chạy
    def run(self):
        wb = openpyxl.load_workbook(self.src, read_only=True, data_only=True)
        ws = find_sheet(wb, "DATA base", "DATAbase", "DATA")
        if ws is None:
            raise SystemExit('Không thấy tab "DATA base" trong file. Các tab có: ' + ", ".join(wb.sheetnames))

        suppliers = self.read_catalog(find_sheet(wb, "listnhacungcap", "list nha cung cap"), SUPPLIER_HEADERS, "name")
        projects = self.read_catalog(find_sheet(wb, "QL Ma DA/HM", "QL Ma DA HM", "QL Ma DAHM", "QL Mã DA/HM"), PROJECT_HEADERS, ("item_code", "project_code"))
        groups = self.read_material_groups(find_sheet(wb, "Nhóm VT", "Nhom VT"))

        rows = read_rows(ws)
        hi, cols, unknown = detect_header(rows, DATA_HEADERS, "name")
        if hi is None:
            raise SystemExit('Tab "DATA base": không tìm thấy dòng tiêu đề có cột "Tên Vật tư".')
        header = rows[hi]
        if unknown:
            self.notes.append("Cột không nhận diện (bỏ qua): " + ", ".join(unknown))
        missing = [f for f in ("qty_requested", "request_code", "request_date", "unit_price") if f not in cols]
        if missing:
            self.notes.append("CẢNH BÁO thiếu cột: " + ", ".join(missing))

        # Cột "Nhà cc/Người mua" là người/bộ phận mua (KHVT, Mr. Trường…), không phải mã NCC → không đối chiếu listnhacungcap
        self.notes.append('Cột "Nhà cc/Người mua" giữ nguyên, không đối chiếu với listnhacungcap (là người mua, không phải mã NCC).')
        project_keys = {fold(p.get("item_code") or p.get("project_code")) for p in projects if p.get("item_code") or p.get("project_code")}
        code_by_fold = {fold(g["name"]): g["name"] for g in groups}
        group_lookup = dict(code_by_fold)
        for g in groups:
            for seg in (g.get("description") or "").split(";"):
                if fold(seg):
                    group_lookup.setdefault(fold(seg), g["name"])
        for alias, target in LEGACY_GROUP_ALIASES.items():
            if fold(target) in code_by_fold:
                group_lookup.setdefault(alias, code_by_fold[fold(target)])

        def cell(r, f):
            idx = cols.get(f)
            return r[idx[0]] if idx and idx[0] < len(r) else None

        def col_name(f):
            return str(header[cols[f][0]]) if f in cols else f

        parsed = []
        total_sheet_amount = 0.0
        for i, r in enumerate(rows[hi + 1:], start=hi + 2):
            data_fields = [f for f in cols if f not in ("stt", "status", "days_left", "week", "month", "qty_missing", "amount")]
            if all(is_empty(cell(r, f)) for f in data_fields):
                continue  # dòng trống
            line = {"_row": i}
            for f in ("unit", "material_group", "name", "uom", "project_code", "project_name", "delivery_location",
                      "supplier", "requester", "equipment_name"):
                line[f] = clean_text(cell(r, f))
            purposes = [clean_text(r[c]) for c in cols.get("purpose", []) if c < len(r)]
            line["purpose"] = next((p for p in purposes if p), None)
            work = clean_text(cell(r, "work_code"))
            line["note"] = f"PL mã CV: {work}" if work else None

            for f in ("qty_requested", "qty_on_hand", "unit_price"):
                v, e = parse_num(cell(r, f))
                if e:
                    self.err(i, col_name(f), cell(r, f), e)
                line[f] = v
            amt, e = parse_num(cell(r, "amount"))
            if amt:
                total_sheet_amount += amt
                calc = (line["qty_requested"] or 0) * (line["unit_price"] or 0)
                if abs(calc - amt) >= 1:
                    self.err(i, col_name("amount"), cell(r, "amount"), f"Thành tiền trong Sheet khác SL × Đơn giá ({calc:,.0f})")

            need_hint, _ = parse_date(cell(r, "need_date"))  # để suy năm khi Ngày yêu cầu ghi kiểu '11-Thg3'
            req_date, e = parse_date(cell(r, "request_date"), need_hint)
            if e:
                self.err(i, col_name("request_date"), cell(r, "request_date"), e)
            line["request_date"] = req_date
            for f, target in (("need_date", "need_date"), ("received_date", "_received_date")):
                d, e = parse_date(cell(r, f), req_date)
                if e:
                    self.err(i, col_name(f), cell(r, f), e)
                line[target] = d
            q, e = parse_num(cell(r, "qty_received"))
            if e:
                self.err(i, col_name("qty_received"), cell(r, "qty_received"), e)
            line["_qty_received"] = q
            line["_signed_by"] = clean_text(cell(r, "signed_by"))  # thực tế ghi "Đủ"/"thiếu", không phải tên người
            line["_sheet_status"] = clean_text(cell(r, "status"))
            line["_raw_code"] = clean_text(cell(r, "request_code"))

            if not line["name"]:
                self.err(i, col_name("name"), None, "Thiếu tên vật tư")
            if line["project_code"] and project_keys and fold(line["project_code"]) not in project_keys:
                self.err(i, col_name("project_code"), line["project_code"], "Không có trong danh mục QL Ma DA/HM")
            if line["material_group"]:
                mapped_group = group_lookup.get(fold(line["material_group"]))
                if mapped_group:
                    line["material_group"] = mapped_group
                elif group_lookup:
                    self.err(i, col_name("material_group"), line["material_group"], "Không có trong danh mục Nhóm VT")
            if line["_received_date"] and line["_qty_received"] is None:
                self.err(i, col_name("qty_received"), None, "Có Ngày nhận nhưng thiếu Số lượng thực nhận (không tạo lần nhận)")
            parsed.append(line)

        self.build(parsed, suppliers, projects, groups, total_sheet_amount)
        return self

    def build(self, parsed, suppliers, projects, groups, total_sheet_amount):
        dates = [p["request_date"] for p in parsed if p["request_date"]]
        newest = max(dates) if dates else dt.date.today()
        pending_from = newest - dt.timedelta(days=self.pending_days)

        # Gán số phiếu: hợp lệ giữ nguyên; sai mẫu / trống (cũ) → phiếu lịch sử LS-0001…
        legacy_keys: OrderedDict[tuple, str] = OrderedDict()
        self.pending_count = 0
        for p in parsed:
            raw = (p["_raw_code"] or "").replace(" ", "").upper()
            if CODE_RE.match(raw):
                p["request_code"], p["legacy_flag"] = raw, False
                continue
            if not raw and (p["request_date"] is None or p["request_date"] >= pending_from):
                p["request_code"], p["legacy_flag"] = None, False  # dòng mới chưa làm phiếu → vẫn chờ xuất phiếu trong app
                self.pending_count += 1
                continue
            key = (p["_raw_code"] or "", iso(p["request_date"]) or "", p["requester"] or "", p["unit"] or "")
            if key not in legacy_keys:
                legacy_keys[key] = f"LS-{len(legacy_keys) + 1:04d}"
            p["request_code"], p["legacy_flag"] = legacy_keys[key], True
            if p["_raw_code"]:
                self.err(p["_row"], "Số phiếu", p["_raw_code"], f"Không đúng mẫu YCyyyy-n → gán phiếu lịch sử {p['request_code']}")

        # lines + receipts
        self.lines, self.receipts = [], []
        for idx, p in enumerate(parsed, start=1):
            created = f"{iso(p['request_date']) or newest.isoformat()}T00:00:00"
            self.lines.append({
                "id": idx, "request_code": p["request_code"], "unit": p["unit"], "material_group": p["material_group"],
                "name": p["name"], "uom": p["uom"], "qty_requested": num_out(p["qty_requested"]),
                "qty_on_hand": num_out(p["qty_on_hand"]), "need_date": iso(p["need_date"]),
                "project_code": p["project_code"], "project_name": p["project_name"], "purpose": p["purpose"],
                "equipment_name": p["equipment_name"], "delivery_location": p["delivery_location"],
                "requester": p["requester"], "request_date": iso(p["request_date"]), "supplier": p["supplier"],
                "unit_price": num_out(p["unit_price"]), "note": p["note"], "legacy_flag": p["legacy_flag"],
                "created_by": "import", "created_at": created, "updated_at": created,
            })
            if p["_qty_received"] is not None and p["_qty_received"] != 0:
                self.receipts.append({
                    "id": len(self.receipts) + 1, "line_id": idx, "received_date": iso(p["_received_date"]),
                    "qty": num_out(p["_qty_received"]), "received_by": None,
                    "note": "import" + (f"; Ký nhận: {p['_signed_by']}" if p["_signed_by"] else ""),
                    "created_at": created,
                })

        # requests: gom theo số phiếu
        by_code: OrderedDict[str, list] = OrderedDict()
        for line in self.lines:
            if line["request_code"]:
                by_code.setdefault(line["request_code"], []).append(line)
        self.requests = []
        for code, ls in by_code.items():
            m = CODE_RE.match(code)
            first_date = min((l["request_date"] for l in ls if l["request_date"]), default=None)
            year = int(m[1]) if m else int((first_date or newest.isoformat())[:4])
            seq = int(m[2]) if m else int(code[3:])

            def join(f):
                return "; ".join(dict.fromkeys(str(l[f]) for l in ls if l[f]))
            need = sorted(l["need_date"] for l in ls if l["need_date"])
            self.requests.append({
                "id": len(self.requests) + 1, "code": code, "year": year, "seq": seq,
                "issued_at": f"{first_date}T00:00:00" if first_date else None, "issued_by": ls[0]["requester"],
                "header_json": json.dumps({
                    "unit": join("unit"), "project_code": join("project_code"), "project_name": join("project_name"),
                    "delivery_location": join("delivery_location"), "requester": join("requester"),
                    "request_date": first_date, "need_date": need[0] if need else None, "note": "",
                }, ensure_ascii=False),
                "line_count": len(ls), "legacy_flag": not m,
            })

        # bộ đếm số phiếu theo năm = số lớn nhất đã dùng (không nhỏ hơn mốc cấu hình)
        max_seq = defaultdict(int)
        for r in self.requests:
            if not r["legacy_flag"]:
                max_seq[r["year"]] = max(max_seq[r["year"]], r["seq"])
        for y, s in INITIAL_SEQ.items():
            max_seq[y] = max(max_seq[y], s)
        self.sequences = [{"id": i, "key": f"YC{y}", "last_value": s} for i, (y, s) in enumerate(sorted(max_seq.items()), 1)]

        def named(values):
            return [{"id": i, "name": v} for i, v in enumerate(values, 1)]

        def distinct(field):
            # gộp "cái/Cái/CÁI " thành 1 mục danh mục (lấy cách viết phổ biến nhất); dữ liệu dòng giữ nguyên
            counts = defaultdict(lambda: defaultdict(int))
            for l in self.lines:
                if l[field]:
                    counts[fold(l[field])][l[field]] += 1
            return sorted((max(v.items(), key=lambda kv: kv[1])[0] for v in counts.values()), key=fold)
        fallback_groups = [{"name": v, "description": None} for v in distinct("material_group")]
        self.catalogs = {
            "suppliers": [dict(id=i, **{k: s.get(k) for k in ("code", "name", "address", "phone", "email", "tax_code", "field")}) for i, s in enumerate(suppliers, 1)],
            "project_items": [dict(id=i, **{k: p.get(k) for k in ("project_code", "project_name", "item_code", "item_name", "owner", "status")}) for i, p in enumerate(projects, 1)],
            "material_groups": [{"id": i, "name": g["name"], "description": g.get("description")} for i, g in enumerate(groups or fallback_groups, 1)],
            "units": named(distinct("unit")),
            "uoms": named(distinct("uom")),
        }
        self.total_sheet_amount = total_sheet_amount

        # Đối chiếu cột "Trạng thái" của Sheet với quy tắc của app: ngày nhận − ngày cần có
        cmp = defaultdict(int)
        self.status_examples = []
        for p in parsed:
            st = fold(p["_sheet_status"])
            if not p["_received_date"] or not p["need_date"]:
                cmp["chưa nhận / thiếu ngày (Sheet để trống)" if not st else "chưa nhận nhưng Sheet có trạng thái"] += 1
                continue
            late = (p["_received_date"] - p["need_date"]).days
            expect = f"Trễ {late} ngày" if late > 0 else "đúng hẹn"
            ok = st == fold(expect)
            partial = (p["_qty_received"] or 0) < (p["qty_requested"] or 0)
            cmp[("khớp" if ok else "lệch") + (" (nhận chưa đủ SL)" if partial else " (nhận đủ SL)")] += 1
            if not ok and len(self.status_examples) < 5:
                self.status_examples.append(f"dòng {p['_row']}: Sheet '{p['_sheet_status']}' ≠ tính '{expect}'")
        self.status_cmp = dict(cmp)

    # ---- báo cáo
    def report(self) -> str:
        total_amount = sum((l["qty_requested"] or 0) * (l["unit_price"] or 0) for l in self.lines)
        out = [f"BÁO CÁO ĐỐI SOÁT — {self.src.name}", "=" * 60]
        out.append(f"Số dòng vật tư (có dữ liệu)     : {len(self.lines):,}")
        out.append(f"  - có số phiếu YC hợp lệ       : {sum(1 for l in self.lines if l['request_code'] and not l['legacy_flag']):,}")
        out.append(f"  - gán phiếu lịch sử LS-xxxx   : {sum(1 for l in self.lines if l['legacy_flag']):,}")
        out.append(f"  - chưa có phiếu (≤{self.pending_days} ngày gần nhất): {self.pending_count:,}")
        out.append(f"Tổng thành tiền (SL × Đơn giá)  : {total_amount:,.0f}")
        out.append(f"Tổng cột 'Thành tiền' trong Sheet: {self.total_sheet_amount:,.0f}"
                   + ("   ✓ khớp" if abs(total_amount - self.total_sheet_amount) < 1 else "   ✗ LỆCH " + f"{total_amount - self.total_sheet_amount:,.0f}"))
        out.append(f"Số phiếu: {len(self.requests):,} (YC: {sum(1 for r in self.requests if not r['legacy_flag']):,}, LS: {sum(1 for r in self.requests if r['legacy_flag']):,})")
        out.append(f"Lần nhận hàng (receipts)        : {len(self.receipts):,}")
        out.append("Bộ đếm số phiếu: " + ", ".join(f"{s['key']}={s['last_value']}" for s in self.sequences))
        out.append(f"Danh mục: NCC {len(self.catalogs['suppliers'])}, hạng mục {len(self.catalogs['project_items'])}, "
                   f"nhóm VT {len(self.catalogs['material_groups'])}, đơn vị {len(self.catalogs['units'])}, ĐVT {len(self.catalogs['uoms'])}")
        err_rows = len({e["Dòng"] for e in self.errors})
        out.append(f"Số dòng có lỗi/cảnh báo         : {err_rows:,} ({len(self.errors):,} ô) → import_errors.xlsx")
        if self.errors:
            by_reason = pd.DataFrame(self.errors).groupby("Lý do").size().sort_values(ascending=False)
            out += ["", "Lỗi theo loại:"] + [f"  {n:6,}  {reason}" for reason, n in by_reason.items()]
        if self.requests:
            df = pd.DataFrame(self.requests)
            df["Tháng"] = [str(a)[:7] if a else f"{y}-?? (không có ngày yêu cầu)" for a, y in zip(df["issued_at"], df["year"])]
            monthly = df.groupby("Tháng").agg(phieu=("code", "count"), dong=("line_count", "sum"))
            out += ["", "Số phiếu theo tháng:", monthly.to_string()]
        if getattr(self, "status_cmp", None):
            out += ["", "Đối chiếu cột Trạng thái (Sheet) với cách tính 'ngày nhận − ngày cần':"]
            out += [f"  {n:6,}  {k}" for k, n in sorted(self.status_cmp.items(), key=lambda x: -x[1])]
            out += [f"         vd. {e}" for e in self.status_examples]
        if self.notes:
            out += ["", "Ghi chú:"] + [f"  - {n}" for n in self.notes]
        return "\n".join(out)

    # ---- ghi file
    def write(self, out_dir: Path):
        out_dir.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        tables = {"lines": self.lines, "requests": self.requests, "receipts": self.receipts, "sequences": self.sequences, **self.catalogs}
        for name, rows in tables.items():
            ws = wb.create_sheet(name)
            cols = self.schema[name]
            ws.append(cols)
            for c in ws[1]:
                c.font = Font(bold=True)
            for r in rows:
                ws.append([_cell(r.get(c)) for c in cols])
            for row in ws.iter_rows():
                for c in row:
                    c.number_format = "@"
            ws.freeze_panes = "A2"
        wb.save(out_dir / "db_khoi_tao.xlsx")
        pd.DataFrame(self.errors, columns=["Dòng", "Cột", "Giá trị", "Lý do"]).to_excel(out_dir / "import_errors.xlsx", index=False)
        (out_dir / "import_report.txt").write_text(self.report(), encoding="utf-8")


def _cell(v):
    """Mọi ô ghi dạng chữ (cột định dạng @) để Google Sheets không tự đổi kiểu."""
    if v is None:
        return None
    if isinstance(v, bool):
        return "TRUE" if v else None
    return str(v)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Import Google Sheet cũ (.xlsx) sang cấu trúc app Đề nghị cấp vật tư")
    ap.add_argument("xlsx", type=Path, help="file .xlsx tải từ Google Sheet cũ")
    ap.add_argument("--out", type=Path, default=ROOT / "out", help="thư mục kết quả (mặc định out/)")
    ap.add_argument("--dry-run", action="store_true", help="chỉ đọc và in báo cáo, không ghi file")
    ap.add_argument("--pending-days", type=int, default=30,
                    help="dòng không có số phiếu trong N ngày gần nhất được giữ nguyên để xuất phiếu trong app (mặc định 30)")
    args = ap.parse_args(argv)
    if not args.xlsx.exists():
        raise SystemExit(f"Không thấy file {args.xlsx}")
    imp = Importer(args.xlsx, args.pending_days).run()
    if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
        sys.stdout.reconfigure(encoding="utf-8")
    print(imp.report())
    if not args.dry_run:
        imp.write(args.out)
        print(f"\nĐã ghi: {args.out / 'db_khoi_tao.xlsx'}, {args.out / 'import_errors.xlsx'}, {args.out / 'import_report.txt'}")
    return imp


if __name__ == "__main__":
    main()
