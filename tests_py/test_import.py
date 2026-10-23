import datetime as dt
import hashlib
import json
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
import unicodedata

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import import_sheet as imp  # noqa: E402

# Đúng tiêu đề tab "DATA base" trong Sheet thật
HEADER = ["STT (tự động)", "Đơn vị đề xuất", "Nhóm VT", "Tên Vật tư", "ĐVT", "Số lượng yêu cầu", "SL tồn kho", "Ngày cần có",
          "Mã Hạng mục", "Hạng mục", "Mục đích sử dụng", "Mục đích sử dụng2", "Đia điểm giao hàng", "Nhà cc/Người mua",
          "Ngày yêu cầu", "Ngày nhận", "Số lượng thực nhận", "Người yêu cầu", "Ký nhận", "Trạng thái", "Số phiếu",
          "Số ngày còn lại", "Đơn giá", "Thành tiền", "Tên thiết bị/hạng mục", "PL mã CV", "tuần", "Tháng", "Số lượng còn thiếu"]


def row(**kw):
    base = dict.fromkeys(HEADER)
    base.update({"Đơn vị đề xuất": "Xưởng CK", "Nhóm VT": "Thép", "ĐVT": "Cái", "Mã Hạng mục": "HM01",
                 "Hạng mục": "Khung", "Đia điểm giao hàng": "Kho 1", "Nhà cc/Người mua": "KHVT", "Người yêu cầu": "Nam"})
    for k, v in kw.items():
        base[k] = v
    return [base[h] for h in HEADER]


@pytest.fixture
def sheet(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DATA base"
    ws.append(["ĐỀ NGHỊ CẤP VẬT TƯ"])  # dòng tiêu đề phụ phía trên
    ws.append(HEADER)
    rows = [
        # cần 23/12/2021, nhận đủ 24/12/2021 → 1 lần nhận
        row(**{"Tên Vật tư": "Thép tấm", "Số lượng yêu cầu": 5, "Ngày cần có": "23/12/2021", "Ngày yêu cầu": "20/12/2021",
               "Ngày nhận": "24-thg 12", "Số lượng thực nhận": 5, "Ký nhận": "Đủ", "Số phiếu": "YC2021-001",
               "Trạng thái": "Trễ 1 ngày",
               "Đơn giá": 1000, "Thành tiền": 5000, "Mục đích sử dụng": "Làm khung"}),
        row(**{"Tên Vật tư": "Bulong M12", "Số lượng yêu cầu": "1.500", "Ngày cần có": 44560, "Ngày yêu cầu": "20-12-2021",
               "Số phiếu": "YC2021-001", "Đơn giá": "2,5", "Thành tiền": 3750}),
        # số phiếu sai mẫu → LS
        row(**{"Tên Vật tư": "Sơn", "Số lượng yêu cầu": 2, "Ngày yêu cầu": dt.datetime(2022, 3, 1), "Số phiếu": "Nam làm phiếu",
               "Đơn giá": 50000, "Thành tiền": 100000, "Nhóm VT": "Nhóm lạ"}),
        # trống số phiếu, cũ → LS ; "Nhà cc/Người mua" là người mua → giữ nguyên; đơn vị viết kiểu tổ hợp (NFD)
        row(**{"Tên Vật tư": "Dây điện", "Số lượng yêu cầu": 10, "Ngày yêu cầu": "11-Thg3", "Ngày cần có": dt.datetime(2022, 3, 15),
               "Nhà cc/Người mua": "Mr. Trường", "Đơn vị đề xuất": unicodedata.normalize("NFD", "P.Cải tiến")}),
        # trống số phiếu, gần đây → chờ xuất phiếu
        row(**{"Tên Vật tư": "Đá cắt", "Số lượng yêu cầu": 3, "Ngày yêu cầu": "20/10/2026", "Ngày cần có": "05-thg 1",
               "Mục đích sử dụng": "Cắt", }),
        row(**{"Tên Vật tư": "Que hàn", "Số lượng yêu cầu": 1, "Ngày yêu cầu": "21/10/2026", "Số phiếu": "yc2026-294",
               "Ngày cần có": "ngày mai", "Ngày nhận": "22/10/2026"}),
        [None] * len(HEADER),  # dòng trống bị bỏ qua
    ]
    rows[4][HEADER.index("Mục đích sử dụng2")] = "Khung-Cắt"  # cột 2 = "Hạng mục-Mục đích" → bỏ khi cột 1 có
    for r in rows:
        ws.append(r)
    ncc = wb.create_sheet("listnhacungcap")
    ncc.append(["STT", "Mã NCC", "Tên NCC", "Địa chỉ", "SĐT"])
    ncc.append([1, "KHVT", "Phòng KHVT", "Nội bộ", "0123"])
    hm = wb.create_sheet("QL Ma DAHM")  # Google bỏ "/" khi xuất; tab có 2 bảng đặt cạnh nhau
    hm.append(["STT", "Mã Dự án (MDM)", "Tên dự án  (MDM)", "Mã hạng mục  (MDM)", "Hạng mục  (MDM)", "Phụ trách", "Trạng thái", "ghi chú",
               "STT", "Mã Dự án (MDM)", "Tên dự án (MDM)", "Mã hạng mục (MDM)", "Hạng mục (MDM)", "Phụ trách"])
    hm.append([1, "HM01", None, None, "Khung", None, None, None, 2, "HM02", None, None, "Lò NL", None])
    g = wb.create_sheet("Nhóm VT")
    g.append(["Nhóm vật tư", None])
    g.append(["Thép", "Thép các loại"])
    path = tmp_path / "sheet.xlsx"
    wb.save(path)
    return path


def run(path, tmp_path):
    out = tmp_path / "out"
    return imp.main([str(path), "--out", str(out)]), out


def test_parse_date_cac_dang():
    ref = dt.date(2021, 12, 20)
    assert imp.parse_date("23/12/2021")[0] == dt.date(2021, 12, 23)
    assert imp.parse_date("23-12-2021")[0] == dt.date(2021, 12, 23)
    assert imp.parse_date(44553)[0] == dt.date(2021, 12, 23)
    assert imp.parse_date("24-thg 12", ref)[0] == dt.date(2021, 12, 24)
    assert imp.parse_date("05-thg 1", ref)[0] == dt.date(2022, 1, 5)  # sang năm sau
    assert imp.parse_date("24/12", ref)[0] == dt.date(2021, 12, 24)
    assert imp.parse_date("24-thg 12")[1]  # thiếu năm, không có ngày yêu cầu → lỗi
    assert imp.parse_date("31/02/2021")[1] == "Ngày không tồn tại"
    assert imp.parse_date("abc")[1] == "Không hiểu ngày"


def test_parse_num():
    assert imp.parse_num("1.234,5")[0] == 1234.5
    assert imp.parse_num("1,5")[0] == 1.5
    assert imp.parse_num("1.500") == (1500, "Cảnh báo: '1.500' được hiểu là 1500 (dấu chấm phân cách nghìn) — kiểm tra lại")
    assert imp.parse_num("0.500")[0] == 0.5
    assert imp.parse_num("2.5")[0] == 2.5
    assert imp.parse_num("1.500.000")[0] == 1500000
    assert imp.parse_num("2 bộ")[1] == "Không phải số"


def test_import_day_du(sheet, tmp_path):
    result, out = run(sheet, tmp_path)
    lines = {l["name"]: l for l in result.lines}
    assert len(result.lines) == 6  # dòng trống bị bỏ
    assert lines["Thép tấm"]["request_code"] == "YC2021-001"
    assert lines["Thép tấm"]["need_date"] == "2021-12-23"
    assert lines["Bulong M12"]["need_date"] == "2021-12-30"  # serial Excel
    assert lines["Bulong M12"]["unit_price"] == 2.5
    assert lines["Sơn"]["request_code"] == "LS-0001" and lines["Sơn"]["legacy_flag"] is True
    assert lines["Dây điện"]["request_code"] == "LS-0002"
    assert lines["Đá cắt"]["request_code"] is None  # gần đây → chờ xuất phiếu trong app
    assert lines["Đá cắt"]["need_date"] == "2027-01-05"
    assert lines["Đá cắt"]["purpose"] == "Cắt"  # cột 1 có → bỏ cột "Hạng mục-Mục đích"
    assert lines["Dây điện"]["request_date"] == "2022-03-11"  # '11-Thg3' lấy năm theo Ngày cần có
    assert lines["Dây điện"]["unit"] == "P.Cải tiến"  # chuẩn hóa NFC
    assert {p["item_name"] for p in result.catalogs["project_items"]} == {"Khung", "Lò NL"}  # đọc cả 2 bảng
    assert lines["Que hàn"]["request_code"] == "YC2026-294"  # chữ thường được chuẩn hóa

    # nhận hàng: ngày thiếu năm lấy năm từ Ngày yêu cầu
    rec = [r for r in result.receipts if r["line_id"] == lines["Thép tấm"]["id"]]
    assert rec == [dict(rec[0], received_date="2021-12-24", qty=5, received_by=None, note="import; Ký nhận: Đủ")]
    assert result.status_cmp.get("khớp (nhận đủ SL)") == 1  # 'Trễ 1 ngày' trùng cách tính của app

    reqs = {r["code"]: r for r in result.requests}
    assert reqs["YC2021-001"]["line_count"] == 2 and reqs["YC2021-001"]["year"] == 2021 and reqs["YC2021-001"]["seq"] == 1
    assert json.loads(reqs["YC2021-001"]["header_json"])["unit"] == "Xưởng CK"
    assert reqs["LS-0001"]["legacy_flag"] is True
    assert {s["key"]: s["last_value"] for s in result.sequences} == {"YC2021": 1, "YC2026": 294}

    # đối soát tổng tiền: 5*1000 + 1500*2.5 + 2*50000 = 108750 = tổng cột Thành tiền
    assert result.total_sheet_amount == 108750
    assert "✓ khớp" in result.report()

    errors = pd.read_excel(out / "import_errors.xlsx")
    reasons = set(errors["Lý do"])
    assert "Không có trong danh mục Nhóm VT" in reasons
    assert "Có Ngày nhận nhưng thiếu Số lượng thực nhận (không tạo lần nhận)" in reasons
    assert "Không hiểu ngày" in reasons  # "ngày mai"
    assert any(r.startswith("Không đúng mẫu") for r in reasons)
    # giá trị lạ được giữ nguyên, không tự sửa
    assert lines["Dây điện"]["supplier"] == "Mr. Trường"


def test_file_ket_qua_dung_cau_truc(sheet, tmp_path):
    _, out = run(sheet, tmp_path)
    wb = openpyxl.load_workbook(out / "db_khoi_tao.xlsx")
    schema = imp.load_schema()
    for tab in ("lines", "requests", "receipts", "sequences", "suppliers", "project_items", "material_groups", "units", "uoms"):
        assert [c.value for c in wb[tab][1]] == schema[tab]
    assert wb["lines"].cell(2, 1).number_format == "@"
    assert wb["suppliers"].cell(2, 2).value == "KHVT"


def test_chay_lai_khong_nhan_doi(sheet, tmp_path):
    _, out = run(sheet, tmp_path)
    first = hashlib.md5(pd.read_excel(out / "db_khoi_tao.xlsx", sheet_name=None)["lines"].to_csv().encode()).hexdigest()
    _, out = run(sheet, tmp_path)
    data = pd.read_excel(out / "db_khoi_tao.xlsx", sheet_name=None)
    assert len(data["lines"]) == 6
    assert hashlib.md5(data["lines"].to_csv().encode()).hexdigest() == first


def test_dry_run_khong_ghi_file(sheet, tmp_path):
    out = tmp_path / "dry"
    imp.main([str(sheet), "--out", str(out), "--dry-run"])
    assert not out.exists()


def test_nguon_khong_bi_sua(sheet, tmp_path):
    before = hashlib.md5(sheet.read_bytes()).hexdigest()
    run(sheet, tmp_path)
    assert hashlib.md5(sheet.read_bytes()).hexdigest() == before
